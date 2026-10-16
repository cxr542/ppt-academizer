"""Geometry-only layout QA for converted academy decks (no rendering needed).

Flags, on content slides only:
  HEADER_OVERLAP   a picture or large shape crosses the bottom of the header band
  PLACEHOLDER_HIT  a non-placeholder shape overlaps the title / governing / step-number box
  NUMBER_WRAP      a multi-digit step number sits in a box too narrow for its digits

These are the defects that used to be found only by opening the deck in PowerPoint (see
CHANGELOG 1.7.13-1.7.15). Usage::

    python -m core.layout_qa converted.pptx [more.pptx ...]   # exit code 1 when issues are found
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .ppt_test_path import ensure_engine_on_path

ensure_engine_on_path()

# Ignore slivers: overlaps smaller than ~0.02in x 0.02in.
_MIN_HIT_AREA = 20_000 * 20_000
# A shape this large (in EMU^2) counts as "large" even when it is not a picture.
_LARGE_AREA = 600_000 * 600_000
# Tolerance above the header bottom before a crossing counts (decorative lines hug the edge).
_HEADER_TOLERANCE = 150_000
# Width one digit of the step number needs (same constant the builder uses).
_EMU_PER_DIGIT = 170_000


@dataclass(frozen=True)
class Issue:
    slide: int
    kind: str
    message: str

    def __str__(self) -> str:
        return f"slide {self.slide:>2}  {self.kind:<15} {self.message}"


def _box(shape) -> tuple[int, int, int, int]:
    left, top = int(shape.left or 0), int(shape.top or 0)
    return left, top, left + int(shape.width or 0), top + int(shape.height or 0)


def _intersection(a, b) -> int:
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0


def check_presentation(path: str | Path) -> list[Issue]:
    from scripts.shape_migrate_lib import CONTENT_LAYOUTS, HEADER_BOTTOM_EMU

    prs = Presentation(str(path))
    slide_width = int(prs.slide_width)
    issues: list[Issue] = []
    for index, slide in enumerate(prs.slides, 1):
        if slide.slide_layout.name not in CONTENT_LAYOUTS:
            continue
        header_boxes = [
            shape
            for shape in slide.shapes
            if shape.is_placeholder
            and shape.has_text_frame
            and shape.text_frame.text.strip()
            and int(shape.top or 0) < HEADER_BOTTOM_EMU + 400_000
            and int(shape.width or 0) < slide_width * 0.95
        ]
        for shape in slide.shapes:
            if shape.is_placeholder:
                text = shape.text_frame.text.strip() if shape.has_text_frame else ""
                if text.isdigit() and len(text) > 1 and int(shape.width) < _EMU_PER_DIGIT * len(text) * 0.9:
                    issues.append(Issue(index, "NUMBER_WRAP", f"'{text}' in a {int(shape.width) / 914400:.2f}in box"))
                continue
            box = _box(shape)
            if box[2] <= box[0] or box[3] <= box[1]:
                continue
            large = (box[2] - box[0]) * (box[3] - box[1]) > _LARGE_AREA
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            if box[1] < HEADER_BOTTOM_EMU - _HEADER_TOLERANCE and box[3] > HEADER_BOTTOM_EMU and (large or is_picture):
                issues.append(
                    Issue(
                        index,
                        "HEADER_OVERLAP",
                        f"{shape.name[:24]} spans {box[1] / 914400:.2f}in-{box[3] / 914400:.2f}in",
                    )
                )
            for header in header_boxes:
                if _intersection(box, _box(header)) > _MIN_HIT_AREA:
                    issues.append(
                        Issue(index, "PLACEHOLDER_HIT", f"{shape.name[:24]} over '{header.text_frame.text.strip()[:20]}'")
                    )
                    break
    return issues


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("decks", nargs="+", type=Path)
    args = parser.parse_args(argv)
    total = 0
    for deck in args.decks:
        issues = check_presentation(deck)
        total += len(issues)
        print(f"== {deck.name}: {len(issues)} issue(s)")
        for issue in issues:
            print(f"   {issue}")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
