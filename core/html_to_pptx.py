"""Static HTML document -> plain .pptx (input stage of the academize pipeline).

The output is an ordinary 16:9 deck (title placeholder + text boxes/tables/pictures) so the
existing academize pipeline restyles it exactly like an uploaded .pptx. Only the source text,
tables and pictures are restructured; nothing is invented.

Slide rules (generic, not sample specific):
  h1            -> cover slide (title, plus the eyebrow/subtitle paragraphs around it)
  h2            -> content slide; when it has h3 children it becomes a divider slide
                   and every h3 gets its own content slide
  overflow      -> content is split over several slides, tables repeat their header row
Scripts, styles, navigation and buttons are ignored; JS-rendered content is not executed.
"""

from __future__ import annotations

import base64
import io
import re
import unicodedata
import urllib.parse
from dataclasses import dataclass, field
from pathlib import Path

from lxml import html as lxml_html
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN_IN = 0.6
BODY_TOP_IN = 1.45
BODY_BOTTOM_IN = 7.0
BODY_W_IN = SLIDE_W_IN - 2 * MARGIN_IN
BODY_H_IN = BODY_BOTTOM_IN - BODY_TOP_IN

TEXT_PT = 16
CODE_PT = 11
NOTE_PT = 14
TABLE_PT = 12
LINE_SPACING = 1.35
BLOCK_GAP_IN = 0.14
MONO_FONT = "Consolas"

_SKIP_TAGS = {"script", "style", "nav", "footer", "head", "button", "noscript", "template", "form", "iframe"}
_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}
_NOTE_CLASS = re.compile(r"\b(note|callout|alert|warning|tip|admonition|notice|hint)\b", re.I)
_META_CLASS = re.compile(r"\b(eyebrow|subtitle|lead|kicker)\b", re.I)
_MAX_IMAGE_BYTES = 12 * 1024 * 1024

Runs = list[tuple[str, bool, bool]]  # (text, bold, code)


@dataclass
class Block:
    kind: str  # p | meta | note | list | table | code | image | subhead
    runs: Runs = field(default_factory=list)
    items: list[tuple[int, Runs]] = field(default_factory=list)  # list: (depth, runs)
    ordered: bool = False
    rows: list[list[str]] = field(default_factory=list)  # table
    header_rows: int = 0
    text: str = ""  # code
    image: bytes = b""
    image_size: tuple[int, int] = (0, 0)


@dataclass
class Section:
    level: int
    title: str
    blocks: list[Block] = field(default_factory=list)
    children: list["Section"] = field(default_factory=list)


def html_to_pptx(src: Path, dst: Path, *, warnings: list[str] | None = None) -> Path:
    """Convert ``src`` (.html/.htm) to a plain .pptx at ``dst``."""
    warns = warnings if warnings is not None else []
    raw = Path(src).read_bytes()
    doc = lxml_html.document_fromstring(_decode(raw))
    title_tag = doc.findtext(".//title") or ""
    body = doc.find("body")
    if body is None:
        body = doc
    root = _build_tree(body, warns)
    deck_title = _clean(root.title) or _clean(title_tag) or Path(src).stem
    if not root.title:
        root.title = deck_title
    prs = _render(root, deck_title)
    prs.save(str(dst))
    return dst


# ---------------------------------------------------------------- parsing


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _classes(el) -> str:
    return el.get("class", "") or ""


def _runs(el, bold: bool = False, code: bool = False) -> Runs:
    """Flatten inline content of ``el`` into (text, bold, code) runs."""
    out: Runs = []

    def add(text: str | None, b: bool, c: bool) -> None:
        if text:
            out.append((re.sub(r"\s+", " ", text) if not c else text, b, c))

    def walk(node, b: bool, c: bool) -> None:
        add(node.text, b, c)
        for child in node:
            if not isinstance(child.tag, str):
                add(child.tail, b, c)
                continue
            tag = child.tag.lower()
            if tag in _SKIP_TAGS or tag in ("img", "svg", "ul", "ol", "table", "pre"):
                add(child.tail, b, c)
                continue
            if tag == "br":
                add(" ", b, c)
            else:
                walk(child, b or tag in ("b", "strong", "th"), c or tag in ("code", "kbd", "samp", "tt"))
            add(child.tail, b, c)

    walk(el, bold, code)
    return _trim(out)


def _trim(runs: Runs) -> Runs:
    runs = [r for r in runs if r[0]]
    if runs:
        runs[0] = (runs[0][0].lstrip(), runs[0][1], runs[0][2])
        runs[-1] = (runs[-1][0].rstrip(), runs[-1][1], runs[-1][2])
    return [r for r in runs if r[0]]


def _plain(runs: Runs) -> str:
    return "".join(r[0] for r in runs)


def _image_block(data: bytes, warns: list[str]) -> Block | None:
    if not data or len(data) > _MAX_IMAGE_BYTES:
        warns.append("용량이 너무 큰 그림은 건너뛰었습니다.")
        return None
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as im:
            im.load()
            size = im.size
            if im.format not in ("PNG", "JPEG", "GIF", "BMP"):
                buf = io.BytesIO()
                im.convert("RGBA").save(buf, "PNG")
                data = buf.getvalue()
    except Exception:
        warns.append("읽을 수 없는 그림은 건너뛰었습니다.")
        return None
    return Block("image", image=data, image_size=size)


_SVG_ATTRS = {
    a.lower(): a
    for a in (
        "viewBox", "preserveAspectRatio", "markerWidth", "markerHeight", "markerUnits", "refX", "refY",
        "gradientUnits", "gradientTransform", "patternUnits", "patternTransform", "patternContentUnits",
        "stdDeviation", "clipPathUnits", "maskUnits", "maskContentUnits", "textLength", "lengthAdjust",
        "startOffset", "spreadMethod", "filterUnits", "primitiveUnits", "baseFrequency", "numOctaves",
    )
}
_SVG_TAGS = {
    t.lower(): t
    for t in ("clipPath", "linearGradient", "radialGradient", "textPath", "foreignObject", "feGaussianBlur", "feOffset", "feBlend", "feColorMatrix", "feMerge", "feMergeNode", "feFlood", "feComposite")
}


def _svg_markup(el) -> str:
    """Serialize an HTML-parsed <svg>; the HTML parser lowercased camelCase names, so restore them."""
    for node in el.iter():
        if not isinstance(node.tag, str):
            continue
        node.tag = _SVG_TAGS.get(node.tag, node.tag)
        for k in list(node.attrib):
            fixed = _SVG_ATTRS.get(k)
            if fixed and fixed != k:
                node.attrib[fixed] = node.attrib.pop(k)
    svg = etree.tostring(el, encoding="unicode", method="xml")
    if "xmlns" not in svg.split(">", 1)[0]:
        svg = svg.replace("<svg", '<svg xmlns="http://www.w3.org/2000/svg"', 1)
    return svg


def _svg_to_png(svg: str) -> bytes | None:
    try:
        import cairosvg  # optional: needs the Cairo system library
    except Exception:
        return None
    try:
        return cairosvg.svg2png(bytestring=svg.encode("utf-8"), scale=2.0, background_color="white")
    except Exception:
        return None


def _data_uri_bytes(src: str) -> tuple[str, bytes] | None:
    if not src.startswith("data:"):
        return None
    head, _, payload = src.partition(",")
    mime = head[5:].split(";")[0].lower()
    try:
        if ";base64" in head:
            return mime, base64.b64decode(payload)
        return mime, urllib.parse.unquote_to_bytes(payload)
    except Exception:
        return None


def _picture(el, warns: list[str]) -> Block | None:
    tag = el.tag.lower()
    if tag == "svg":
        png = _svg_to_png(_svg_markup(el))
        if png is None:
            warns.append("SVG 그림을 이미지로 바꾸지 못해 건너뛰었습니다.")
            return None
        return _image_block(png, warns)
    decoded = _data_uri_bytes((el.get("src") or "").strip())
    if decoded is None:
        warns.append("외부 주소의 그림은 가져오지 않아 건너뛰었습니다.")
        return None
    mime, data = decoded
    if mime == "image/svg+xml":
        png = _svg_to_png(data.decode("utf-8", errors="replace"))
        if png is None:
            warns.append("SVG 그림을 이미지로 바꾸지 못해 건너뛰었습니다.")
            return None
        data = png
    return _image_block(data, warns)


def _table_block(el) -> Block | None:
    rows: list[list[str]] = []
    header_rows = 0
    for tr in el.iter("tr"):
        cells = [c for c in tr if isinstance(c.tag, str) and c.tag.lower() in ("td", "th")]
        if not cells:
            continue
        if not rows and all(c.tag.lower() == "th" for c in cells):
            header_rows = 1
        rows.append([_clean(c.text_content()) for c in cells])
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return Block("table", rows=rows, header_rows=header_rows)


def _list_items(el, depth: int, out: list[tuple[int, Runs]]) -> None:
    for li in el:
        if not isinstance(li.tag, str) or li.tag.lower() != "li":
            continue
        runs = _runs(li)
        if runs:
            out.append((depth, runs))
        for sub in li:
            if isinstance(sub.tag, str) and sub.tag.lower() in ("ul", "ol"):
                _list_items(sub, depth + 1, out)


def _walk(el, blocks: list, warns: list[str]) -> None:
    """Append heading tuples ("h", level, text) and Block objects in document order."""

    def tail_para(node) -> None:
        text = _clean(node.tail)
        if text:
            blocks.append(Block("p", runs=[(text, False, False)]))

    for child in el:
        if not isinstance(child.tag, str):
            tail_para(child)
            continue
        tag = child.tag.lower()
        if tag in _SKIP_TAGS:
            tail_para(child)
            continue
        if tag in _HEADINGS:
            text = _clean(child.text_content())
            if text:
                blocks.append(("h", _HEADINGS[tag], text))
        elif tag in ("ul", "ol"):
            items: list[tuple[int, Runs]] = []
            _list_items(child, 0, items)
            if items:
                blocks.append(Block("list", items=items, ordered=tag == "ol"))
        elif tag == "table":
            tb = _table_block(child)
            if tb:
                blocks.append(tb)
        elif tag == "pre":
            text = child.text_content().strip("\n")
            if text.strip():
                blocks.append(Block("code", text=text))
        elif tag in ("img", "svg"):
            pic = _picture(child, warns)
            if pic:
                blocks.append(pic)
        elif tag in ("p", "figcaption", "blockquote", "dt", "dd", "summary"):
            runs = _runs(child)
            if runs:
                kind = "note" if tag == "blockquote" else ("meta" if _META_CLASS.search(_classes(child)) else "p")
                blocks.append(Block(kind, runs=runs))
            for pic_el in child.iter("img", "svg"):
                pic = _picture(pic_el, warns)
                if pic:
                    blocks.append(pic)
        else:
            if _NOTE_CLASS.search(_classes(child)) and not any(
                isinstance(c.tag, str) and c.tag.lower() in ("table", "pre", "ul", "ol") for c in child
            ):
                runs = _runs(child)
                if runs:
                    blocks.append(Block("note", runs=runs))
            elif len(child) == 0 or all(
                isinstance(c.tag, str) and c.tag.lower() in ("b", "strong", "i", "em", "code", "a", "span", "br", "small", "kbd")
                for c in child
            ):
                runs = _runs(child)
                if runs:
                    kind = "meta" if _META_CLASS.search(_classes(child)) else "p"
                    blocks.append(Block(kind, runs=runs))
            else:
                lead = _clean(child.text)
                if lead:
                    blocks.append(Block("p", runs=[(lead, False, False)]))
                _walk(child, blocks, warns)
        tail_para(child)


def _build_tree(body, warns: list[str]) -> Section:
    flat: list = []
    _walk(body, flat, warns)
    root = Section(1, "")
    stack = [root]
    seen_h1 = False
    for item in flat:
        if isinstance(item, tuple):
            _, level, text = item
            if level == 1 and not seen_h1 and not root.children:
                root.title = text
                seen_h1 = True
                continue
            if level >= 4:
                stack[-1].blocks.append(Block("subhead", runs=[(text, True, False)]))
                continue
            sec = Section(max(level, 2), text)
            while len(stack) > 1 and stack[-1].level >= sec.level:
                stack.pop()
            stack[-1].children.append(sec)
            stack.append(sec)
        else:
            stack[-1].blocks.append(item)
    return root


# ---------------------------------------------------------------- layout estimate


def _char_w_em(ch: str) -> float:
    return 1.0 if unicodedata.east_asian_width(ch) in ("W", "F") else 0.55


def _lines(text: str, width_in: float, pt: float) -> int:
    cap = max(width_in * 72.0, 1.0)
    total = 0
    for part in text.split("\n") or [""]:
        used = 0.0
        n = 1
        for ch in part:
            w = _char_w_em(ch) * pt
            if used + w > cap:
                n += 1
                used = 0.0
            used += w
        total += n
    return total


def _line_h(pt: float) -> float:
    return pt * LINE_SPACING / 72.0


def _text_h(text: str, width_in: float, pt: float) -> float:
    return _lines(text, width_in, pt) * _line_h(pt)


def _col_widths(rows: list[list[str]], total_in: float) -> list[float]:
    n = len(rows[0])
    weights = []
    for c in range(n):
        longest = max((sum(_char_w_em(ch) for ch in r[c]) for r in rows), default=1)
        weights.append(min(max(longest, 4.0), 40.0))
    s = sum(weights)
    return [total_in * w / s for w in weights]


def _row_h(row: list[str], widths: list[float]) -> float:
    lines = max(_lines(cell, max(w - 0.2, 0.4), TABLE_PT) for cell, w in zip(row, widths))
    return lines * _line_h(TABLE_PT) + 0.14


def _block_h(b: Block) -> float:
    if b.kind in ("p", "meta"):
        return _text_h(_plain(b.runs), BODY_W_IN, TEXT_PT) + BLOCK_GAP_IN
    if b.kind == "subhead":
        return _text_h(_plain(b.runs), BODY_W_IN, TEXT_PT) + BLOCK_GAP_IN
    if b.kind == "note":
        return _text_h(_plain(b.runs), BODY_W_IN - 0.4, NOTE_PT) + 0.3 + BLOCK_GAP_IN
    if b.kind == "list":
        return sum(_text_h(_plain(r), BODY_W_IN - 0.35 - 0.3 * d, TEXT_PT) + 0.04 for d, r in b.items) + BLOCK_GAP_IN
    if b.kind == "code":
        return _text_h(b.text, BODY_W_IN - 0.4, CODE_PT) + 0.3 + BLOCK_GAP_IN
    if b.kind == "table":
        widths = _col_widths(b.rows, BODY_W_IN)
        return sum(_row_h(r, widths) for r in b.rows) + BLOCK_GAP_IN
    if b.kind == "image":
        return _image_box(b, BODY_W_IN, BODY_H_IN)[1] + BLOCK_GAP_IN
    return BLOCK_GAP_IN


def _image_box(b: Block, max_w: float, max_h: float) -> tuple[float, float]:
    w, h = b.image_size
    if not w or not h:
        return max_w, max_h
    scale = min(1.0, max_w / (w / 96.0), max_h / (h / 96.0))
    return (w / 96.0) * scale, (h / 96.0) * scale


def _split_block(b: Block, room: float) -> tuple[Block | None, Block | None]:
    """Split an oversized block so the first part fits ``room`` inches (None if unsplittable)."""
    if b.kind == "table":
        widths = _col_widths(b.rows, BODY_W_IN)
        head = b.rows[: b.header_rows]
        used = sum(_row_h(r, widths) for r in head)
        body = b.rows[b.header_rows :]
        take = 0
        for r in body:
            h = _row_h(r, widths)
            if used + h > room:
                break
            used += h
            take += 1
        if take == 0 or take == len(body):
            return None, None
        first = Block("table", rows=head + body[:take], header_rows=b.header_rows)
        rest = Block("table", rows=head + body[take:], header_rows=b.header_rows)
        return first, rest
    if b.kind == "list":
        used, take = 0.0, 0
        for d, r in b.items:
            h = _text_h(_plain(r), BODY_W_IN - 0.35 - 0.3 * d, TEXT_PT) + 0.04
            if used + h > room:
                break
            used += h
            take += 1
        if take == 0 or take == len(b.items):
            return None, None
        return (
            Block("list", items=b.items[:take], ordered=b.ordered),
            Block("list", items=b.items[take:], ordered=b.ordered),
        )
    if b.kind == "code":
        lines = b.text.split("\n")
        per = max(int((room - 0.3) / _line_h(CODE_PT)), 0)
        if per < 2 or per >= len(lines):
            return None, None
        return Block("code", text="\n".join(lines[:per])), Block("code", text="\n".join(lines[per:]))
    return None, None


def _paginate(blocks: list[Block]) -> list[list[Block]]:
    pages: list[list[Block]] = []
    cur: list[Block] = []
    used = 0.0
    queue = list(blocks)
    while queue:
        b = queue.pop(0)
        h = _block_h(b)
        if used + h <= BODY_H_IN or (not cur and h <= BODY_H_IN * 1.0):
            cur.append(b)
            used += h
            continue
        if cur and h > BODY_H_IN * 0.35:
            first, rest = _split_block(b, BODY_H_IN - used)
            if first is not None:
                cur.append(first)
                queue.insert(0, rest)
                pages.append(cur)
                cur, used = [], 0.0
                continue
        if cur:
            pages.append(cur)
            cur, used = [], 0.0
            queue.insert(0, b)
            continue
        # empty page and still too tall
        first, rest = _split_block(b, BODY_H_IN)
        if first is not None:
            pages.append([first])
            queue.insert(0, rest)
        else:
            pages.append([b])  # image/paragraph taller than a page: scaled to fit when drawn
    if cur:
        pages.append(cur)
    return pages


# ---------------------------------------------------------------- rendering


def _render(root: Section, deck_title: str) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    subtitle = " · ".join(_plain(b.runs) for b in root.blocks if b.kind == "meta")
    _cover(prs, deck_title, subtitle)
    lead = [b for b in root.blocks if b.kind != "meta"]
    if lead:
        _content(prs, deck_title, lead)
    for sec in root.children:
        _emit_section(prs, sec)
    return prs


def _emit_section(prs: Presentation, sec: Section) -> None:
    subs = [c for c in sec.children]
    if subs:
        _divider(prs, sec.title)
        if sec.blocks:
            _content(prs, sec.title, sec.blocks)
        for c in subs:
            _emit_section(prs, c)
    else:
        _content(prs, sec.title, sec.blocks)


def _title_shape(slide, text: str, *, size: int, top: float, height: float, bold: bool = True) -> None:
    ph = slide.shapes.title
    ph.left, ph.top = Inches(MARGIN_IN), Inches(top)
    ph.width, ph.height = Inches(BODY_W_IN), Inches(height)
    ph.text_frame.word_wrap = True
    ph.text_frame.text = text
    for p in ph.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold


def _cover(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    _title_shape(slide, title, size=40, top=2.4, height=1.6)
    if subtitle:
        _textbox(slide, [[(subtitle, False, False)]], MARGIN_IN, 4.2, BODY_W_IN, 1.0, 20)


def _divider(prs: Presentation, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _title_shape(slide, title, size=36, top=2.8, height=1.6)


def _content(prs: Presentation, title: str, blocks: list[Block]) -> None:
    pages = _paginate(blocks) or [[]]
    for i, page in enumerate(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        label = title if len(pages) == 1 else f"{title} ({i + 1}/{len(pages)})"
        _title_shape(slide, label, size=28, top=0.4, height=0.9)
        y = BODY_TOP_IN
        for b in page:
            y += _draw(slide, b, y)


def _set_runs(par, runs: Runs, size: float, *, color: RGBColor | None = None) -> None:
    for text, bold, code in runs:
        r = par.add_run()
        r.text = text
        r.font.size = Pt(size if not code else max(size - 1, 9))
        r.font.bold = bold or None
        if code:
            r.font.name = MONO_FONT
        if color is not None:
            r.font.color.rgb = color


def _textbox(slide, paragraphs: list[Runs], x: float, y: float, w: float, h: float, size: float):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, runs in enumerate(paragraphs):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_runs(par, runs, size)
    return tb


def _draw(slide, b: Block, y: float) -> float:
    """Draw ``b`` at ``y`` and return the vertical space it used."""
    h = _block_h(b)
    room = max(BODY_BOTTOM_IN - y, 0.5)
    if b.kind in ("p", "meta", "subhead"):
        _textbox(slide, [b.runs], MARGIN_IN, y, BODY_W_IN, h - BLOCK_GAP_IN, TEXT_PT)
    elif b.kind == "list":
        tb = slide.shapes.add_textbox(Inches(MARGIN_IN), Inches(y), Inches(BODY_W_IN), Inches(h - BLOCK_GAP_IN))
        tf = tb.text_frame
        tf.word_wrap = True
        counters: dict[int, int] = {}
        for i, (depth, runs) in enumerate(b.items):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.level = min(depth, 4)
            marker = "•"
            if b.ordered:
                counters[depth] = counters.get(depth, 0) + 1
                marker = f"{counters[depth]}."
            _set_runs(par, [(f"{marker} ", False, False)] + runs, TEXT_PT)
            par.space_after = Pt(2)
    elif b.kind == "note":
        box = slide.shapes.add_shape(1, Inches(MARGIN_IN), Inches(y), Inches(BODY_W_IN), Inches(h - BLOCK_GAP_IN))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xFF, 0xF8, 0xE1)
        box.line.color.rgb = RGBColor(0xE0, 0xA8, 0x00)
        tf = box.text_frame
        tf.word_wrap = True
        _set_runs(tf.paragraphs[0], b.runs, NOTE_PT, color=RGBColor(0x1F, 0x23, 0x28))
        tf.paragraphs[0].alignment = 1
    elif b.kind == "code":
        box = slide.shapes.add_shape(1, Inches(MARGIN_IN), Inches(y), Inches(BODY_W_IN), Inches(h - BLOCK_GAP_IN))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF6)
        box.line.color.rgb = RGBColor(0xD9, 0xDD, 0xE1)
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(b.text.split("\n")):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = 1
            r = par.add_run()
            r.text = line
            r.font.size = Pt(CODE_PT)
            r.font.name = MONO_FONT
            r.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    elif b.kind == "table":
        widths = _col_widths(b.rows, BODY_W_IN)
        heights = [_row_h(r, widths) for r in b.rows]
        shape = slide.shapes.add_table(
            len(b.rows), len(b.rows[0]), Inches(MARGIN_IN), Inches(y), Inches(BODY_W_IN), Inches(sum(heights))
        )
        tbl = shape.table
        tbl.first_row = bool(b.header_rows)
        for c, w in enumerate(widths):
            tbl.columns[c].width = Inches(w)
        for r, row in enumerate(b.rows):
            tbl.rows[r].height = Inches(heights[r])
            for c, text in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text_frame.word_wrap = True
                par = cell.text_frame.paragraphs[0]
                run = par.add_run()
                run.text = text
                run.font.size = Pt(TABLE_PT)
    elif b.kind == "image":
        w, ih = _image_box(b, BODY_W_IN, min(BODY_H_IN, room))
        x = MARGIN_IN + (BODY_W_IN - w) / 2
        slide.shapes.add_picture(io.BytesIO(b.image), Inches(x), Inches(y), Inches(w), Inches(ih))
        h = ih + BLOCK_GAP_IN
    return h
