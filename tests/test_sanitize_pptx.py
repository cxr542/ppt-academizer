"""Tests for pptx package finalization (PowerPoint repair dialog prevention)."""

from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation

from scripts.sanitize_pptx_package import (
    _drop_parts_from_content_types,
    _drop_rels,
    _scrub_presentation_xml,
    finalize_pptx_package,
)


def test_drop_parts_keeps_types_root() -> None:
    ct = (
        '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/ppt/fonts/font1.fntdata" ContentType="application/x-fontdata"/>'
        '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        "</Types>"
    )
    out = _drop_parts_from_content_types(ct)
    assert "<Types" in out
    assert "/ppt/fonts/" not in out
    assert "/ppt/slides/slide1.xml" in out


def test_drop_rels_scrubs_embedded_font_refs() -> None:
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId10" Type="font" Target="fonts/font1.fntdata"/>'
        '<Relationship Id="rId22" Type="slide" Target="slides/slide1.xml"/>'
        "</Relationships>"
    )
    pres = (
        '<p:presentation embedTrueTypeFonts="1">'
        "<p:embeddedFontLst><p:embeddedFont><p:regular r:id=\"rId10\"/></p:embeddedFont></p:embeddedFontLst>"
        '<p:sldIdLst><p:sldId r:id="rId22"/></p:sldIdLst></p:presentation>'
    )
    cleaned, dropped = _drop_rels(rels)
    assert "rId10" not in cleaned
    assert "rId22" in cleaned
    assert dropped == {"rId10"}
    out = _scrub_presentation_xml(pres, dropped)
    assert "embeddedFontLst" not in out
    assert 'r:id="rId10"' not in out
    assert 'r:id="rId22"' in out
    assert 'embedTrueTypeFonts="0"' in out


def test_finalize_preserves_openable_package(tmp_path: Path) -> None:
    src = tmp_path / "mini.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(src))

    finalize_pptx_package(src)

    with zipfile.ZipFile(src) as z:
        ct = z.read("[Content_Types].xml").decode("utf-8")
    assert "<Types" in ct
    assert len(ct) > 100

    prs2 = Presentation(str(src))
    assert len(prs2.slides) >= 1


def test_finalize_after_second_pass_idempotent(tmp_path: Path) -> None:
    src = tmp_path / "twice.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(src))
    finalize_pptx_package(src)
    ct1 = zipfile.ZipFile(src).read("[Content_Types].xml").decode("utf-8")
    finalize_pptx_package(src)
    ct2 = zipfile.ZipFile(src).read("[Content_Types].xml").decode("utf-8")
    assert "<Types" in ct2
    assert len(ct2) >= len(ct1) * 0.9
