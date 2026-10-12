"""SmartArt / OLE / hyperlinks: cloned shapes must bring their relationships along."""

from __future__ import annotations

import posixpath
import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from scripts.shape_migrate_lib import clone_shape_xml
from scripts.slide_classifier_v2 import analyze_slide_layout, has_data_object

from smartart_sample import build


def _sample(tmp_path: Path):
    path = tmp_path / "sa.pptx"
    build(str(path))
    return Presentation(str(path))


def _dangling(path: Path) -> list[str]:
    """Slide XML relationship attributes that do not resolve in the slide's .rels."""
    z = zipfile.ZipFile(path)
    bad: list[str] = []
    for name in z.namelist():
        if not re.match(r"ppt/slides/slide\d+\.xml$", name):
            continue
        rels = {
            r.get("Id")
            for r in etree.fromstring(z.read(name.replace("slides/", "slides/_rels/") + ".rels"))
        }
        used = set(re.findall(rb'r:(?:id|embed|link|dm|lo|qs|cs)="(rId\d+)"', z.read(name)))
        bad += [f"{name}:{u.decode()}" for u in used if u.decode() not in rels]
    return bad


def _clone_frame_to_new_deck(tmp_path: Path, slide_index: int) -> Path:
    src = _sample(tmp_path)
    frame = [s for s in src.slides[slide_index].shapes if s._element.tag.endswith("}graphicFrame")][0]
    dst = Presentation()
    dst_slide = dst.slides.add_slide(dst.slide_layouts[6])
    clone_shape_xml(frame, dst_slide)
    out = tmp_path / "out.pptx"
    dst.save(out)
    return out


def test_graphic_frames_are_data_objects_not_covers(tmp_path: Path) -> None:
    src = _sample(tmp_path)
    for idx in (1, 2):  # SmartArt slide, OLE slide: a title and a frame python-pptx cannot see into
        assert has_data_object(src.slides[idx])
        assert analyze_slide_layout(src.slides[idx])["layout"] != "cover"
        assert analyze_slide_layout(src.slides[idx])["body_heavy"] is True
    assert not has_data_object(src.slides[0])


def test_cloned_smartart_brings_all_diagram_parts(tmp_path: Path) -> None:
    out = _clone_frame_to_new_deck(tmp_path, 1)
    z = zipfile.ZipFile(out)
    parts = sorted(n for n in z.namelist() if n.startswith("ppt/diagrams/"))
    assert len(parts) == 5, parts  # data, layout, quickStyle, colors, drawing
    assert _dangling(out) == []

    # the data model must still name the drawing part, through the NEW slide rId
    data = z.read("ppt/diagrams/data1.xml").decode("utf-8")
    rid = re.search(r'dataModelExt[^>]*relId="(rId\d+)"', data).group(1)
    rels = {
        r.get("Id"): r.get("Target")
        for r in etree.fromstring(z.read("ppt/slides/_rels/slide1.xml.rels"))
    }
    assert posixpath.basename(rels[rid]) == "drawing1.xml"
    assert "계획" in z.read("ppt/diagrams/drawing1.xml").decode("utf-8")


def test_cloned_ole_object_keeps_embedding_and_preview(tmp_path: Path) -> None:
    out = _clone_frame_to_new_deck(tmp_path, 2)
    z = zipfile.ZipFile(out)
    assert any(n.startswith("ppt/embeddings/") for n in z.namelist())
    assert any(n.startswith("ppt/media/") for n in z.namelist())  # preview picture
    assert _dangling(out) == []


def test_cloned_hyperlink_keeps_its_external_relationship(tmp_path: Path) -> None:
    src = Presentation()
    slide = src.slides.add_slide(src.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "링크"
    run.hyperlink.address = "https://example.com/doc"

    dst = Presentation()
    dst_slide = dst.slides.add_slide(dst.slide_layouts[6])
    clone_shape_xml(box, dst_slide)
    out = tmp_path / "link.pptx"
    dst.save(out)

    z = zipfile.ZipFile(out)
    assert b"https://example.com/doc" in z.read("ppt/slides/_rels/slide1.xml.rels")
    assert _dangling(out) == []


def test_smartart_starting_in_the_header_band_is_body_content(tmp_path: Path) -> None:
    from scripts.shape_migrate_lib import is_body_shape

    src = _sample(tmp_path)
    frame = [s for s in src.slides[1].shapes if s._element.tag.endswith("}graphicFrame")][0]
    frame.top = Inches(0.5)  # inside the header band
    assert is_body_shape(frame, "업무 프로세스", None)
