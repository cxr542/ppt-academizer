"""Synthetic pptx with a SmartArt diagram (hand-made parts) on slide 2 and an embedded-OLE xlsx
on slide 3. python-pptx cannot author SmartArt, so the five diagram parts are written by hand."""
import io
from lxml import etree
from pptx import Presentation
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt
from pptx.enum.shapes import PROG_ID
import xlsxwriter

NS_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_DSP = "http://schemas.microsoft.com/office/drawing/2008/diagram"
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
LABELS = ["계획", "실행", "점검"]

def data_xml(drawing_rid):
    pts, cxns = [], []
    doc = "{00000000-0000-0000-0000-000000000000}"
    pts.append(f'<dgm:pt modelId="{doc}" type="doc"><dgm:prSet loTypeId="urn:microsoft.com/office/officeart/2005/8/layout/default" loCatId="list" qsTypeId="urn:microsoft.com/office/officeart/2005/8/quickstyle/simple1" qsCatId="simple" csTypeId="urn:microsoft.com/office/officeart/2005/8/colors/accent1_2" csCatId="accent1" phldr="0"/><dgm:spPr/><dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="ko-KR"/></a:p></dgm:t></dgm:pt>')
    for i, label in enumerate(LABELS, 1):
        nid, par, sib = (f"{{0000000{i}-0000-0000-0000-00000000000{k}}}" for k in (1, 2, 3))
        pts.append(f'<dgm:pt modelId="{nid}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="ko-KR"/><a:t>{label}</a:t></a:r></a:p></dgm:t></dgm:pt>')
        pts.append(f'<dgm:pt modelId="{par}" type="parTrans" cxnId="{{0000009{i}-0000-0000-0000-000000000000}}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="ko-KR"/></a:p></dgm:t></dgm:pt>')
        pts.append(f'<dgm:pt modelId="{sib}" type="sibTrans" cxnId="{{0000009{i}-0000-0000-0000-000000000000}}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="ko-KR"/></a:p></dgm:t></dgm:pt>')
        cxns.append(f'<dgm:cxn modelId="{{0000009{i}-0000-0000-0000-000000000000}}" srcId="{doc}" destId="{nid}" srcOrd="{i-1}" destOrd="0" parTransId="{par}" sibTransId="{sib}"/>')
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><dgm:dataModel xmlns:dgm="{NS_DGM}" xmlns:a="{NS_A}">'
            f'<dgm:ptLst>{"".join(pts)}</dgm:ptLst><dgm:cxnLst>{"".join(cxns)}</dgm:cxnLst><dgm:bg/><dgm:whole/>'
            f'<dgm:extLst><a:ext uri="http://schemas.microsoft.com/office/drawing/2008/diagram"><dsp:dataModelExt xmlns:dsp="{NS_DSP}" relId="{drawing_rid}" minVer="http://schemas.openxmlformats.org/drawingml/2006/diagram"/></a:ext></dgm:extLst></dgm:dataModel>').encode()

LAYOUT = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><dgm:layoutDef xmlns:dgm="{NS_DGM}" xmlns:a="{NS_A}" uniqueId="urn:microsoft.com/office/officeart/2005/8/layout/default">'
          '<dgm:title val=""/><dgm:desc val=""/><dgm:catLst><dgm:cat type="list" pri="400"/></dgm:catLst>'
          '<dgm:layoutNode name="diagram"><dgm:varLst><dgm:dir/><dgm:resizeHandles val="exact"/></dgm:varLst><dgm:alg type="snake"><dgm:param type="grDir" val="tL"/><dgm:param type="flowDir" val="row"/></dgm:alg>'
          '<dgm:shape xmlns:r="' + NS_R + '" r:blip=""><dgm:adjLst/></dgm:shape><dgm:presOf/>'
          '<dgm:forEach name="Name0" axis="ch" ptType="node"><dgm:layoutNode name="node"><dgm:varLst><dgm:bulletEnabled val="1"/></dgm:varLst><dgm:alg type="tx"/>'
          '<dgm:shape xmlns:r="' + NS_R + '" type="rect" r:blip=""><dgm:adjLst/></dgm:shape><dgm:presOf axis="desOrSelf" ptType="node"/></dgm:layoutNode></dgm:forEach></dgm:layoutNode></dgm:layoutDef>').encode()
QSTYLE = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><dgm:styleDef xmlns:dgm="{NS_DGM}" xmlns:a="{NS_A}" uniqueId="urn:microsoft.com/office/officeart/2005/8/quickstyle/simple1">'
          '<dgm:title val=""/><dgm:desc val=""/><dgm:catLst><dgm:cat type="simple" pri="10100"/></dgm:catLst><dgm:scene3d><a:camera prst="orthographicFront"/><a:lightRig rig="threePt" dir="t"/></dgm:scene3d>'
          '<dgm:styleLbl name="node1"><dgm:scene3d><a:camera prst="orthographicFront"/><a:lightRig rig="threePt" dir="t"/></dgm:scene3d><dgm:sp3d/><dgm:txPr/>'
          '<dgm:style><a:lnRef idx="2"><a:scrgbClr r="0" g="0" b="0"/></a:lnRef><a:fillRef idx="1"><a:scrgbClr r="0" g="0" b="0"/></a:fillRef><a:effectRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:effectRef><a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></dgm:style></dgm:styleLbl></dgm:styleDef>').encode()
COLORS = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><dgm:colorsDef xmlns:dgm="{NS_DGM}" xmlns:a="{NS_A}" uniqueId="urn:microsoft.com/office/officeart/2005/8/colors/accent1_2">'
          '<dgm:title val=""/><dgm:desc val=""/><dgm:catLst><dgm:cat type="accent1" pri="11200"/></dgm:catLst>'
          '<dgm:styleLbl name="node1"><dgm:fillClrLst meth="repeat"><a:schemeClr val="accent1"/></dgm:fillClrLst><dgm:linClrLst meth="repeat"><a:schemeClr val="lt1"/></dgm:linClrLst><dgm:effectClrLst/><dgm:txLinClrLst/><dgm:txFillClrLst/><dgm:txEffectClrLst/></dgm:styleLbl></dgm:colorsDef>').encode()

def drawing_xml():
    sps = []
    for i, label in enumerate(LABELS):
        x = 200000 + i * 3200000
        sps.append(f'<dsp:sp modelId="{{0000000{i+1}-0000-0000-0000-000000000001}}"><dsp:nvSpPr><dsp:cNvPr id="0" name=""/><dsp:cNvSpPr/></dsp:nvSpPr><dsp:spPr><a:xfrm><a:off x="{x}" y="400000"/><a:ext cx="3000000" cy="1800000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></dsp:spPr><dsp:style><a:lnRef idx="2"><a:scrgbClr r="0" g="0" b="0"/></a:lnRef><a:fillRef idx="1"><a:scrgbClr r="0" g="0" b="0"/></a:fillRef><a:effectRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:effectRef><a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></dsp:style><dsp:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="ko-KR" sz="2800"/><a:t>{label}</a:t></a:r></a:p></dsp:txBody><dsp:txXfrm><a:off x="{x}" y="400000"/><a:ext cx="3000000" cy="1800000"/></dsp:txXfrm></dsp:sp>')
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><dsp:drawing xmlns:dgm="{NS_DGM}" xmlns:dsp="{NS_DSP}" xmlns:a="{NS_A}"><dsp:spTree><dsp:nvGrpSpPr><dsp:cNvPr id="0" name=""/><dsp:cNvGrpSpPr/></dsp:nvGrpSpPr><dsp:grpSpPr/>{"".join(sps)}</dsp:spTree></dsp:drawing>').encode()

def build(path):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    def title(s, text, pt=40):
        b = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1.2)); b.text_frame.text = text
        b.text_frame.paragraphs[0].runs[0].font.size = Pt(pt)
    s0 = prs.slides.add_slide(blank); title(s0, "SmartArt·OLE 재현 과정", 44)

    s1 = prs.slides.add_slide(blank); title(s1, "업무 프로세스")
    pkg = s1.part.package
    mk = lambda tmpl, ct, blob: Part(pkg.next_partname(tmpl), ct, pkg, blob)
    drawing = mk("/ppt/diagrams/drawing%d.xml", "application/vnd.ms-office.drawingml.diagramDrawing+xml", drawing_xml())
    # relate the drawing first so its slide rId can be written into the data model
    rid_dr = s1.part.relate_to(drawing, "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing")
    data = mk("/ppt/diagrams/data%d.xml", "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml", data_xml(rid_dr))
    layout = mk("/ppt/diagrams/layout%d.xml", "application/vnd.openxmlformats-officedocument.drawingml.diagramLayout+xml", LAYOUT)
    qs = mk("/ppt/diagrams/quickStyle%d.xml", "application/vnd.openxmlformats-officedocument.drawingml.diagramStyle+xml", QSTYLE)
    cs = mk("/ppt/diagrams/colors%d.xml", "application/vnd.openxmlformats-officedocument.drawingml.diagramColors+xml", COLORS)
    rdm, rlo, rqs, rcs = (s1.part.relate_to(p, RT + t) for p, t in ((data, "diagramData"), (layout, "diagramLayout"), (qs, "diagramQuickStyle"), (cs, "diagramColors")))
    frame = etree.fromstring(
        f'<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{NS_A}" xmlns:r="{NS_R}">'
        '<p:nvGraphicFramePr><p:cNvPr id="20" name="Diagram 19"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="914400" y="1800000"/><a:ext cx="10058400" cy="2600000"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{NS_DGM}"><dgm:relIds xmlns:dgm="{NS_DGM}" r:dm="{rdm}" r:lo="{rlo}" r:qs="{rqs}" r:cs="{rcs}"/></a:graphicData></a:graphic></p:graphicFrame>')
    s1.shapes._spTree.append(frame)

    s2 = prs.slides.add_slide(blank); title(s2, "첨부 문서(OLE)")
    xl = io.BytesIO(); wb = xlsxwriter.Workbook(xl, {"in_memory": True}); ws = wb.add_worksheet(); ws.write(0, 0, "항목"); ws.write(0, 1, "값"); ws.write(1, 0, "수강생"); ws.write(1, 1, 120); wb.close(); xl.seek(0)
    s2.shapes.add_ole_object(xl, PROG_ID.XLSX, Inches(1), Inches(2), Inches(3), Inches(2))
    prs.save(path)
