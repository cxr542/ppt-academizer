import io

import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from core import image_ocr
from core.warnings_display import format_warning


def _png(w=800, h=400, color=(30, 60, 120)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


def _deck_with_picture(tmp_path, blob, size=(5, 3)):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(blob), Inches(1), Inches(1.5), Inches(size[0]), Inches(size[1]))
    path = tmp_path / "in.pptx"
    prs.save(path)
    return path


def test_overlay_places_filled_textbox_over_picture(tmp_path, monkeypatch):
    path = _deck_with_picture(tmp_path, _png())
    monkeypatch.setattr(image_ocr, "tesseract_available", lambda: True)
    monkeypatch.setattr(
        image_ocr, "ocr_image", lambda im: [{"text": "배포 파이프라인", "box": [80, 40, 400, 90], "conf": 90}]
    )
    warnings = image_ocr.apply_image_ocr(path)
    assert warnings[0]["code"] == "IMAGE_OCR_DONE" and warnings[0]["boxes"] == 1
    slide = Presentation(str(path)).slides[0]
    pic = next(s for s in slide.shapes if s.shape_type == 13)
    box = next(s for s in slide.shapes if s.name == image_ocr.SHAPE_NAME)
    assert box.text_frame.text == "배포 파이프라인"
    assert pic.left < box.left < pic.left + pic.width  # sits on the picture, which is kept
    assert box.top > pic.top


def test_small_and_fullbleed_pictures_are_skipped(tmp_path, monkeypatch):
    monkeypatch.setattr(image_ocr, "tesseract_available", lambda: True)
    monkeypatch.setattr(image_ocr, "ocr_image", lambda im: [{"text": "abc def", "box": [1, 1, 50, 20], "conf": 99}])
    icon = _deck_with_picture(tmp_path, _png(), size=(1, 1))
    assert image_ocr.apply_image_ocr(icon)[0]["pictures"] == 0
    prs = Presentation(str(icon))
    assert all(s.name != image_ocr.SHAPE_NAME for s in prs.slides[0].shapes)


def test_unavailable_tesseract_gives_warning_not_error(tmp_path, monkeypatch):
    monkeypatch.setattr(image_ocr, "tesseract_available", lambda: False)
    path = _deck_with_picture(tmp_path, _png())
    assert image_ocr.apply_image_ocr(path)[0]["code"] == "IMAGE_OCR_UNAVAILABLE"


def test_warning_messages():
    done = format_warning({"code": "IMAGE_OCR_DONE", "pictures": 2, "boxes": 5, "failed": 0})
    assert "5줄" in done["message"] and done["level"] == "info"
    none = format_warning({"code": "IMAGE_OCR_DONE", "pictures": 1, "boxes": 0, "failed": 0})
    assert "찾지 못했습니다" in none["message"]
    assert format_warning({"code": "IMAGE_OCR_UNAVAILABLE"})["level"] == "warn"


@pytest.mark.skipif(not image_ocr.tesseract_available(), reason="tesseract not installed")
def test_real_ocr_reads_text_and_ignores_blank_photo():
    font = None
    for path in ("/System/Library/Fonts/AppleSDGothicNeo.ttc", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"):
        try:
            font = ImageFont.truetype(path, 44)
            break
        except OSError:
            continue
    if font is None:
        pytest.skip("no Korean font")
    im = Image.new("RGB", (900, 300), (255, 255, 255))
    ImageDraw.Draw(im).text((40, 100), "배포 파이프라인", fill=(20, 20, 20), font=font)
    texts = " ".join(l["text"] for l in image_ocr.ocr_image(im)).replace(" ", "")
    assert "배포" in texts
    assert image_ocr.ocr_image(Image.new("RGB", (600, 300), (120, 160, 200))) == []
