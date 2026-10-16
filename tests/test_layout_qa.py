"""The layout QA flags the three defects that used to need PowerPoint to spot."""

import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches

import scripts.shape_migrate_lib as sml
from core import layout_qa
from core.layout_qa import check_presentation, main


@pytest.fixture(autouse=True)
def _blank_is_a_content_layout(monkeypatch):
    monkeypatch.setattr(sml, "CONTENT_LAYOUTS", set(sml.CONTENT_LAYOUTS) | {"Blank"})


def _deck(tmp_path, build):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(slide)
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def _png() -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (60, 60), (200, 200, 200)).save(buf, "PNG")
    buf.seek(0)
    return buf


def _kinds(path):
    return {issue.kind for issue in check_presentation(path)}


def test_picture_over_the_header_is_flagged(tmp_path) -> None:
    path = _deck(tmp_path, lambda s: s.shapes.add_picture(_png(), Inches(6.7), Inches(0.5), Inches(5.8), Inches(6.5)))
    assert "HEADER_OVERLAP" in _kinds(path)


def test_picture_below_the_header_is_clean(tmp_path) -> None:
    path = _deck(tmp_path, lambda s: s.shapes.add_picture(_png(), Inches(6.7), Inches(1.6), Inches(5.0), Inches(4.0)))
    assert check_presentation(path) == []


def _title_only_deck(tmp_path, build):
    """A deck whose slide has a real title placeholder (layout "Title Only")."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    build(slide)
    path = tmp_path / "title.pptx"
    prs.save(path)
    return path


@pytest.fixture
def title_only_is_content(monkeypatch):
    monkeypatch.setattr(sml, "CONTENT_LAYOUTS", set(sml.CONTENT_LAYOUTS) | {"Title Only"})


def test_bar_over_a_header_placeholder_is_flagged(tmp_path, title_only_is_content) -> None:
    def build(slide):
        title = slide.shapes.title
        title.text = "제목"
        title.left, title.top, title.width, title.height = Inches(0.85), Inches(0.41), Inches(3), Inches(0.4)
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.35), Inches(0.5), Inches(0.5))

    issues = check_presentation(_title_only_deck(tmp_path, build))
    assert [i.kind for i in issues] == ["PLACEHOLDER_HIT"]


def test_a_shape_clear_of_the_placeholder_is_not_flagged(tmp_path, title_only_is_content) -> None:
    def build(slide):
        title = slide.shapes.title
        title.text = "제목"
        title.left, title.top, title.width, title.height = Inches(0.85), Inches(0.41), Inches(3), Inches(0.4)
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(3), Inches(2), Inches(1))

    assert check_presentation(_title_only_deck(tmp_path, build)) == []


def test_two_digit_number_in_a_one_digit_box_is_flagged(tmp_path, title_only_is_content) -> None:
    def build(slide):
        number = slide.shapes.title
        number.text = "10"
        number.left, number.top, number.width, number.height = Inches(0.56), Inches(0.39), Emu(121_828), Inches(0.24)

    assert [i.kind for i in check_presentation(_title_only_deck(tmp_path, build))] == ["NUMBER_WRAP"]


def test_two_digit_number_in_a_wide_enough_box_is_clean(tmp_path, title_only_is_content) -> None:
    def build(slide):
        number = slide.shapes.title
        number.text = "10"
        number.left, number.top, number.width, number.height = Inches(0.56), Inches(0.39), Inches(0.37), Inches(0.24)

    assert check_presentation(_title_only_deck(tmp_path, build)) == []


def test_step_number_wrap_rule_uses_the_builders_digit_width() -> None:
    from scripts.academy_deck_build_lib import STEP_NUMBER_EMU_PER_DIGIT

    assert layout_qa._EMU_PER_DIGIT == STEP_NUMBER_EMU_PER_DIGIT


def test_header_band_constant_matches_the_engine() -> None:
    from scripts.academy_deck_build_lib import HEADER_BOTTOM_EMU as lib_value

    assert lib_value == sml.HEADER_BOTTOM_EMU


def test_cli_exit_code_reflects_issues(tmp_path, capsys) -> None:
    bad = _deck(tmp_path, lambda s: s.shapes.add_picture(_png(), Inches(6.7), Inches(0.5), Inches(5.8), Inches(6.5)))
    assert main([str(bad)]) == 1
    assert "HEADER_OVERLAP" in capsys.readouterr().out
    good = tmp_path / "good.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(good)
    assert main([str(good)]) == 0


def test_non_content_layouts_are_skipped(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(sml, "CONTENT_LAYOUTS", {"내지_거버닝 O"})
    path = _deck(tmp_path, lambda s: s.shapes.add_picture(_png(), Inches(6.7), Inches(0.5), Inches(5.8), Inches(6.5)))
    assert check_presentation(path) == []
