"""The layout footer shows the year of the conversion, whether the template carries it as a
picture (older templates) or as text (newer ones)."""

import datetime
import io
import json

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches

import scripts.build_cmp_academy as bca
from scripts.deck_migrate_config import footer_copyright_year


class _Layout:
    """A slide stands in for a layout: it has the same shapes/part surface the code uses."""

    def __init__(self, slide):
        self.name, self.shapes, self.part = "내지_거버닝 O", slide.shapes, slide.part


class _Deck:
    def __init__(self, slide, prs):
        self.slide_layouts = [_Layout(slide)]
        self.slide_width, self.slide_height = prs.slide_width, prs.slide_height


def _strip_png() -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGBA", (1460, 60), (201, 201, 208, 109)).save(buf, "PNG")
    buf.seek(0)
    return buf


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return _Deck(slide, prs), slide


def _footer_picture(slide):
    return slide.shapes.add_picture(_strip_png(), Inches(9.63), Inches(7.12), Inches(3.16), Inches(0.13))


def _texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


def test_picture_footer_becomes_text_with_the_given_year() -> None:
    deck, slide = _deck()
    pic = _footer_picture(slide)
    geometry = (pic.left, pic.top, pic.width, pic.height)

    assert bca.apply_current_year_to_footers(deck, 2031) == 1

    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    box = next(s for s in slide.shapes if s.has_text_frame and "OKESTRO" in s.text_frame.text)
    assert box.text_frame.text == "© 2031. OKESTRO Co., Ltd. All Rights Reserved."
    assert (box.left, box.top, box.width, box.height) == geometry
    xml = box._element.xml
    assert 'srgbClr val="C9C9D0"' in xml and "<a:alpha" in xml  # colour and opacity of the original strip


def test_text_footer_year_is_updated_even_across_runs() -> None:
    deck, slide = _deck()
    box = slide.shapes.add_textbox(Inches(9.7), Inches(7.09), Inches(3.2), Inches(0.2))
    para = box.text_frame.paragraphs[0]
    para.add_run().text = "ⓒ 20"
    para.add_run().text = "25. OKESTRO Co., Ltd. All Rights Reserved."

    assert bca.apply_current_year_to_footers(deck, 2031) == 1
    assert box.text_frame.text == "ⓒ 2031. OKESTRO Co., Ltd. All Rights Reserved."


def test_nothing_changes_when_the_year_is_already_right() -> None:
    deck, slide = _deck()
    box = slide.shapes.add_textbox(Inches(9.7), Inches(7.09), Inches(3.2), Inches(0.2))
    box.text_frame.text = "ⓒ 2031. OKESTRO Co., Ltd. All Rights Reserved."
    assert bca.apply_current_year_to_footers(deck, 2031) == 0


def test_other_pictures_and_texts_are_left_alone() -> None:
    deck, slide = _deck()
    tall = slide.shapes.add_picture(_strip_png(), Inches(9.63), Inches(1.0), Inches(3.16), Inches(0.13))  # not at the bottom
    square = slide.shapes.add_picture(_strip_png(), Inches(9.63), Inches(7.0), Inches(0.5), Inches(0.5))  # not a strip
    note = slide.shapes.add_textbox(Inches(1), Inches(7.1), Inches(3), Inches(0.2))
    note.text_frame.text = "© 2025 someone else"
    assert bca.apply_current_year_to_footers(deck, 2031) == 0
    assert tall in list(slide.shapes) and square in list(slide.shapes) and note.text_frame.text == "© 2025 someone else"


def test_default_is_the_conversion_year_and_config_can_pin_or_disable(tmp_path, monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    assert footer_copyright_year() == datetime.date.today().year
    cfg = tmp_path / "kinds.json"
    monkeypatch.setenv("PPT_ACADEMIZER_DECK_KINDS", str(cfg))
    for value, expected in (("auto", datetime.date.today().year), (2030, 2030), ("2029", 2029), (None, None), ("", None)):
        cfg.write_text(json.dumps({"footer_copyright_year": value}), encoding="utf-8")
        assert footer_copyright_year() == expected, value
    deck, slide = _deck()
    _footer_picture(slide)
    cfg.write_text(json.dumps({"footer_copyright_year": None}), encoding="utf-8")
    assert bca.apply_current_year_to_footers(deck) == 0
    assert [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


class _NamedLayout(_Layout):
    def __init__(self, slide, name):
        super().__init__(slide)
        self.name = name


def _deck_named(name):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    class Deck:
        slide_layouts = [_NamedLayout(slide, name)]
        slide_width, slide_height = prs.slide_width, prs.slide_height

    return Deck(), slide


def test_toc_layout_footer_picture_is_removed(monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    deck, slide = _deck_named("목차")
    _footer_picture(slide)
    assert bca.remove_footers_from_layouts(deck) == 1
    assert not list(slide.shapes)


def test_toc_layout_footer_text_is_removed(monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    deck, slide = _deck_named("목차")
    box = slide.shapes.add_textbox(Inches(9.7), Inches(7.09), Inches(3.2), Inches(0.2))
    box.text_frame.text = "ⓒ 2026. OKESTRO Co., Ltd. All Rights Reserved."
    keep = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(3), Inches(0.4))
    keep.text_frame.text = "CONTENTS"
    assert bca.remove_footers_from_layouts(deck) == 1
    assert _texts(slide) == ["CONTENTS"]


def test_other_layouts_keep_their_footer(monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    deck, slide = _deck_named("내지_거버닝 O")
    _footer_picture(slide)
    assert bca.remove_footers_from_layouts(deck) == 0
    assert [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_config_chooses_which_layouts_have_no_footer(tmp_path, monkeypatch) -> None:
    cfg = tmp_path / "kinds.json"
    monkeypatch.setenv("PPT_ACADEMIZER_DECK_KINDS", str(cfg))
    cfg.write_text(json.dumps({"layouts_without_footer": []}), encoding="utf-8")
    deck, slide = _deck_named("목차")
    _footer_picture(slide)
    assert bca.remove_footers_from_layouts(deck) == 0  # [] keeps every footer
    cfg.write_text(json.dumps({"layouts_without_footer": ["간지", "내지_참고"]}), encoding="utf-8")
    deck2, slide2 = _deck_named("내지_참고")
    _footer_picture(slide2)
    assert bca.remove_footers_from_layouts(deck2) == 1
