"""Deck-kind config is data-driven; Google PNG slides keep their aspect ratio."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from PIL import Image

from scripts.academy_deck_build_lib import apply_background_images
from scripts.build_cmp_academy import classify_slide, collect_section_titles
from scripts.deck_migrate_config import (
    GENERIC_KIND,
    config_for_kind,
    detect_deck_kind,
    skip_shape_names,
)


@pytest.fixture(autouse=True)
def _no_user_config(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setenv("PPT_ACADEMIZER_DECK_KINDS", str(tmp_path / "missing.json"))


def test_builtin_cmp_values_unchanged() -> None:
    cfg = config_for_kind(detect_deck_kind("클라우드 구현기술 CMP.pptx"))
    assert cfg.deck_kind == "cmp"
    assert cfg.skip_src_indices == frozenset({8})
    assert cfg.part_cover_indices == frozenset({0, 13})
    assert list(cfg.toc_slide_range) == list(range(2, 8))
    assert cfg.cover_subtitle == "클라우드 구현기술 (CMP)"


def test_ai_filename_wins_over_other_kinds() -> None:
    assert detect_deck_kind("chatgpt-cmp-course.pptx") == GENERIC_KIND


def test_json_adds_kind_and_overrides_builtin(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    f = tmp_path / "kinds.json"
    f.write_text(
        json.dumps(
            {
                "global_skip_shape_names": ["Google Shape;1;p1"],
                "kinds": {
                    "partner_x": {
                        "filename_patterns": ["partner-x"],
                        "skip_src_indices": [3],
                        "part_cover_indices": [5],
                    },
                    "cmp": {"cover_subtitle": "Custom"},
                },
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setenv("PPT_ACADEMIZER_DECK_KINDS", str(f))

    assert detect_deck_kind("Partner-X intro.pptx") == "partner_x"
    cfg = config_for_kind("partner_x")
    assert cfg.skip_src_indices == frozenset({3})
    assert cfg.part_cover_indices == frozenset({0, 5})  # slide 0 is always a cover
    cmp_cfg = config_for_kind("cmp")
    assert cmp_cfg.cover_subtitle == "Custom"
    assert cmp_cfg.part_cover_indices == frozenset({0, 13})  # untouched fields keep defaults
    assert "Google Shape;1;p1" in skip_shape_names()
    assert "Google Shape;520;p15" in skip_shape_names()


def test_malformed_json_falls_back_to_builtins(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    f = tmp_path / "bad.json"
    f.write_text("{not json", encoding="utf-8")
    monkeypatch.setenv("PPT_ACADEMIZER_DECK_KINDS", str(f))
    assert detect_deck_kind("x.cmp.pptx") == "cmp"


def test_classify_slide_has_no_cmp_default_cover() -> None:
    prs = Presentation()
    for _ in range(14):
        prs.slides.add_slide(prs.slide_layouts[6])
    assert classify_slide(prs.slides[13], 13) != "cover"
    assert classify_slide(prs.slides[0], 0) == "cover"
    # empty part_cover set no longer silently becomes {0, 13}
    assert collect_section_titles(prs, 1, part_cover_indices=frozenset()) == []


def _deck_with_image_slide(tmp_path: Path, size: tuple[int, int]):
    img = tmp_path / "s.png"
    Image.new("RGB", size, "white").save(img)
    prs = Presentation()
    layout = prs.slide_layouts[6]
    layout.name = "1_내지_거버닝 X"
    prs.slides.add_slide(layout)
    return prs, {"background_image": str(img)}


def test_background_image_keeps_aspect_ratio_and_is_centered(tmp_path: Path) -> None:
    prs, spec = _deck_with_image_slide(tmp_path, (1600, 900))  # 16:9 into a box that is not 16:9
    assert apply_background_images(prs, [spec]) == 1
    pic = prs.slides[0].shapes[-1]
    assert abs(pic.width / pic.height - 16 / 9) < 0.01
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    assert pic.left >= 0 and pic.top >= 0
    assert pic.left + pic.width <= sw and pic.top + pic.height <= sh


def test_tall_background_image_is_not_stretched_wide(tmp_path: Path) -> None:
    prs, spec = _deck_with_image_slide(tmp_path, (900, 1600))
    apply_background_images(prs, [spec])
    pic = prs.slides[0].shapes[-1]
    assert abs(pic.width / pic.height - 900 / 1600) < 0.01
