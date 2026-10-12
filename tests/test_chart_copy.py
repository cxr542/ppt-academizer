"""Chart graphicFrames must carry their chart part (and workbook) to the output deck."""

from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from scripts.shape_migrate_lib import copy_shape_hybrid


def _deck_with_chart() -> Presentation:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("S", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), data
    )
    return prs


def test_copied_chart_keeps_data_and_workbook(tmp_path: Path) -> None:
    src = _deck_with_chart()
    dst = Presentation()
    dst_slide = dst.slides.add_slide(dst.slide_layouts[6])

    copy_shape_hybrid(dst_slide, src.slides[0].shapes[0])
    out = tmp_path / "out.pptx"
    dst.save(out)

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
    assert "ppt/charts/chart1.xml" in names
    assert any(n.startswith("ppt/embeddings/") and n.endswith(".xlsx") for n in names)

    shape = Presentation(str(out)).slides[0].shapes[0]
    assert shape.has_chart
    series = shape.chart.plots[0].series[0]
    assert list(series.values) == [1.0, 2.0]
    assert list(shape.chart.plots[0].categories) == ["A", "B"]


def test_same_chart_copied_twice_gets_separate_parts(tmp_path: Path) -> None:
    src = _deck_with_chart()
    dst = Presentation()
    dst_slide = dst.slides.add_slide(dst.slide_layouts[6])

    copy_shape_hybrid(dst_slide, src.slides[0].shapes[0])
    copy_shape_hybrid(dst_slide, src.slides[0].shapes[0])
    out = tmp_path / "out.pptx"
    dst.save(out)

    charts = [s.chart for s in Presentation(str(out)).slides[0].shapes if s.has_chart]
    assert len(charts) == 2
    assert charts[0].part.partname != charts[1].part.partname
    wb = [c.part.chart_workbook.xlsx_part.partname for c in charts]
    assert wb[0] != wb[1]


def test_copied_chart_axis_ids_are_unsigned(tmp_path: Path) -> None:
    """python-pptx emits negative axId; the OOXML validator requires xs:unsignedInt."""
    src = _deck_with_chart()
    dst = Presentation()
    dst_slide = dst.slides.add_slide(dst.slide_layouts[6])
    copy_shape_hybrid(dst_slide, src.slides[0].shapes[0])
    out = tmp_path / "out.pptx"
    dst.save(out)

    chart_xml = Presentation(str(out)).slides[0].shapes[0].chart.part.blob.decode("utf-8")
    import re

    ids = re.findall(r'<c:(?:axId|crossAx) val="([^"]+)"', chart_xml)
    assert ids and all(i.isdigit() for i in ids)
    # PowerPoint treats ids as int32; anything above 2**31-1 triggers its repair dialog.
    assert all(int(i) <= 2**31 - 1 for i in ids)
