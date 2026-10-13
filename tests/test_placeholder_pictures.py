"""Only solid-black / stock-size stubs are placeholders; flat design panels are not."""

import io

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from scripts.build_cmp_academy import is_decorative_shadow_picture, is_stock_placeholder_picture


def _picture(color, size=(560, 400), mode="RGB"):
    buf = io.BytesIO()
    Image.new(mode, size, color).save(buf, "PNG")
    buf.seek(0)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return slide.shapes.add_picture(buf, Inches(1), Inches(1), Inches(4), Inches(3))


def test_solid_black_google_stub_is_a_placeholder() -> None:
    assert is_stock_placeholder_picture(_picture((0, 0, 0)))


def test_solid_navy_panel_is_not_a_placeholder() -> None:
    # real deck: a 1280x600 navy (0,31,63) divider panel was dropped, leaving white text on white
    assert not is_stock_placeholder_picture(_picture((0, 31, 63), size=(1280, 600)))


def test_other_flat_colours_are_not_placeholders() -> None:
    for color in ((255, 255, 255), (200, 160, 80), (30, 30, 30)):
        assert not is_stock_placeholder_picture(_picture(color)), color


def test_faint_shadow_is_decoration_not_a_placeholder() -> None:
    # real deck: near-transparent black PNGs (drop shadows) were reported as empty image slots
    for alpha in (0, 17, 77):
        pic = _picture((0, 0, 0, alpha), mode="RGBA")
        assert is_decorative_shadow_picture(pic), alpha
        assert not is_stock_placeholder_picture(pic), alpha


def test_semi_opaque_dark_overlay_is_neither() -> None:
    pic = _picture((0, 0, 0, 160), mode="RGBA")
    assert not is_decorative_shadow_picture(pic)
    assert not is_stock_placeholder_picture(pic)


def test_opaque_black_rgba_stub_is_still_a_placeholder() -> None:
    assert is_stock_placeholder_picture(_picture((0, 0, 0, 255), mode="RGBA"))


def test_stock_size_grey_placeholder_with_transparent_corners_is_still_found() -> None:
    """The "600 x 400" circle placeholder is RGBA with transparent corners."""
    buf = io.BytesIO()
    im = Image.new("RGBA", (400, 400), (0, 0, 0, 0))
    from PIL import ImageDraw

    ImageDraw.Draw(im).ellipse((0, 0, 399, 399), fill=(221, 221, 221, 255))
    im.save(buf, "PNG")
    buf.seek(0)
    prs = Presentation()
    pic = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(buf, Inches(1), Inches(1), Inches(3), Inches(3))
    assert is_stock_placeholder_picture(pic)
