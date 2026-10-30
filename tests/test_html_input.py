"""HTML input: a static HTML document becomes a plain .pptx that the normal pipeline accepts."""

from __future__ import annotations

import asyncio
import base64
import io

import pytest
from fastapi import HTTPException
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from starlette.datastructures import UploadFile

from api.upload_handlers import remove_upload_tmp, require_pptx_filename, save_upload_bounded
from core.html_to_pptx import html_to_pptx


def _png_data_uri() -> str:
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "#3366cc").save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


DOC = f"""<!doctype html><html><head><title>샘플 가이드</title>
<style>p{{color:red}}</style><script>alert(1)</script></head><body>
<nav><a href="#a">목차 링크</a></nav>
<p class="eyebrow">VER 1.0</p><h1>샘플 가이드</h1><p class="subtitle">부제 문장</p>
<h2>1. 개요</h2><p>첫 문단 <b>굵게</b> <code>kubectl get</code></p>
<ul><li>항목 하나<ul><li>하위 항목</li></ul></li><li>항목 둘</li></ul>
<h2>2. 상세</h2>
<h3>2.1 표</h3>
<table><tr><th>이름</th><th>값</th></tr><tr><td>a</td><td>1</td></tr></table>
<div class="note"><b>참고</b>주의 문장</div>
<h3>2.2 코드와 그림</h3>
<pre><code>line1
line2</code></pre>
<img src="{_png_data_uri()}" alt="그림">
<img src="https://example.com/remote.png">
<button>복사</button>
</body></html>"""


def _texts(prs) -> str:
    out = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out.extend(c.text for row in sh.table.rows for c in row.cells)
    return "\n".join(out)


@pytest.fixture
def converted(tmp_path):
    src = tmp_path / "guide.html"
    src.write_text(DOC, encoding="utf-8")
    warns: list[str] = []
    dst = html_to_pptx(src, tmp_path / "guide.pptx", warnings=warns)
    return Presentation(str(dst)), warns


def test_heading_rules_make_cover_divider_and_content_slides(converted):
    prs, _ = converted
    titles = [s.shapes.title.text_frame.text for s in prs.slides]
    # cover, "1. 개요" (no h3 -> one content slide), divider "2. 상세", then one slide per h3
    assert titles == ["샘플 가이드", "1. 개요", "2. 상세", "2.1 표", "2.2 코드와 그림"]


def test_source_text_tables_and_pictures_are_kept_and_noise_is_dropped(converted):
    prs, warns = converted
    text = _texts(prs)
    for expected in ("부제 문장", "첫 문단", "kubectl get", "항목 하나", "하위 항목", "항목 둘", "이름", "참고", "주의 문장", "line1", "line2"):
        assert expected in text
    for noise in ("alert(1)", "목차 링크", "복사", "color:red"):
        assert noise not in text
    content = list(prs.slides)[-1]
    assert sum(1 for sh in content.shapes if sh.shape_type == 13) == 1  # the data: PNG, not the remote one
    assert any("외부 주소" in w for w in warns)
    table_slide = list(prs.slides)[3]
    assert any(getattr(sh, "has_table", False) and sh.has_table for sh in table_slide.shapes)


def test_everything_stays_inside_the_slide(converted):
    prs, _ = converted
    for slide in prs.slides:
        for sh in slide.shapes:
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= prs.slide_width + Inches(0.01)


def test_long_table_is_split_with_repeated_header(tmp_path):
    rows = "".join(f"<tr><td>행{i}</td><td>설명 문장 {i}</td></tr>" for i in range(60))
    src = tmp_path / "long.html"
    src.write_text(f"<h1>긴 표</h1><h2>데이터</h2><table><tr><th>키</th><th>값</th></tr>{rows}</table>", encoding="utf-8")
    prs = Presentation(str(html_to_pptx(src, tmp_path / "long.pptx")))
    pages = [s for s in prs.slides if any(getattr(sh, "has_table", False) and sh.has_table for sh in s.shapes)]
    assert len(pages) > 1
    for s in pages:
        tbl = next(sh.table for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table)
        assert tbl.cell(0, 0).text == "키"
    assert "(1/" in pages[0].shapes.title.text_frame.text
    total = sum(len(next(sh.table for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table).rows) - 1 for s in pages)
    assert total == 60


def test_document_without_headings_still_converts(tmp_path):
    src = tmp_path / "plain.html"
    src.write_text("<html><head><title>제목만</title></head><body><p>본문 한 줄</p></body></html>", encoding="utf-8")
    prs = Presentation(str(html_to_pptx(src, tmp_path / "plain.pptx")))
    assert "본문 한 줄" in _texts(prs)


def test_svg_is_rasterized_when_cairosvg_is_available(tmp_path):
    pytest.importorskip("cairosvg")
    src = tmp_path / "svg.html"
    src.write_text(
        '<h1>도식</h1><h2>구성</h2><svg viewBox="0 0 100 50"><rect x="5" y="5" width="90" height="40" fill="#0a0"/></svg>',
        encoding="utf-8",
    )
    prs = Presentation(str(html_to_pptx(src, tmp_path / "svg.pptx")))
    assert any(sh.shape_type == 13 for s in prs.slides for sh in s.shapes)


def test_upload_accepts_html_and_hands_the_pipeline_a_pptx(monkeypatch):
    monkeypatch.setattr("api.upload_handlers.max_upload_bytes", lambda: 1_000_000)
    upload = UploadFile(filename="가이드.html", file=io.BytesIO(DOC.encode("utf-8")))
    tmp, src = asyncio.run(save_upload_bounded(upload))
    try:
        assert src.name == "가이드.pptx"
        assert len(Presentation(str(src)).slides) >= 5
    finally:
        remove_upload_tmp(tmp)


def test_unsupported_extension_message_names_both_formats():
    require_pptx_filename("a.html")
    require_pptx_filename("a.HTM")
    with pytest.raises(HTTPException) as exc:
        require_pptx_filename("a.docx")
    assert exc.value.status_code == 400 and "HTML" in exc.value.detail
