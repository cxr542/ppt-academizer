"""A thank-you slide is built as a divider (간지) without a chapter number, never as a cover."""

import json

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from scripts.build_cmp_academy import (
    build_slide_plan,
    closing_title,
    is_closing_slide,
)
from scripts.deck_migrate_config import closing_slide_patterns, config_for_kind


def _slide(prs, texts):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for text, pt, top in texts:
        box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(10), Inches(1))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(pt)
    return slide


def _deck(last_texts):
    prs = Presentation()
    _slide(prs, [("강의 제목", 44, 2.0)])  # cover
    _slide(prs, [("본문 제목", 28, 0.6), ("핵심 설명입니다.", 16, 2.0)])
    _slide(prs, last_texts)
    return prs


CLOSING = [("THANK YOU", 15, 0.8), ("경청해 주셔서 감사합니다", 60, 1.4), ("팀 이름", 15, 5.0)]


def test_last_slide_with_thanks_becomes_a_divider_not_a_cover() -> None:
    prs = _deck(CLOSING)
    plan = build_slide_plan(prs)
    assert [k for _, k in plan] == ["cover", "content", "section"]


def test_closing_title_is_the_biggest_line_not_the_kicker() -> None:
    prs = _deck(CLOSING)
    assert closing_title(prs.slides[2]) == "경청해 주셔서 감사합니다"


def test_closing_slide_need_not_be_last_and_the_first_slide_never_is() -> None:
    """Real deck: the thank-you slide is followed by an "Image Sources" slide."""
    prs = Presentation()
    first = _slide(prs, CLOSING)
    _slide(prs, [("본문 제목", 28, 0.6), ("핵심 설명입니다.", 16, 2.0)])
    thanks = _slide(prs, CLOSING)
    _slide(prs, [("Image Sources", 33, 0.6), ("https://example.com/a.jpg", 10, 4.0)])
    cfg = config_for_kind("generic_partner")
    assert [k for _, k in build_slide_plan(prs)] == ["cover", "content", "section", "content"]
    assert is_closing_slide(thanks, 2, len(prs.slides), cfg)
    assert not is_closing_slide(first, 0, len(prs.slides), cfg)


def test_a_small_thanks_line_under_another_title_does_not_count() -> None:
    prs = _deck([("다음 단계 안내", 54, 2.0), ("감사합니다", 14, 4.5)])
    cfg = config_for_kind("generic_partner")
    assert not is_closing_slide(prs.slides[2], 2, 3, cfg)


def test_a_content_slide_mentioning_thanks_is_not_a_closing_slide() -> None:
    long_body = "감사합니다 라는 표현이 들어 있지만 실제로는 긴 본문 설명이 이어지는 내용 슬라이드입니다. " * 2
    prs = _deck([("본문 제목", 28, 0.6), (long_body, 14, 2.0)])
    cfg = config_for_kind("generic_partner")
    assert not is_closing_slide(prs.slides[2], 2, 3, cfg)


def test_patterns_come_from_the_config_file(tmp_path, monkeypatch) -> None:
    cfg_file = tmp_path / "kinds.json"
    cfg_file.write_text(json.dumps({"closing_slide_patterns": ["수고하셨습니다"]}), encoding="utf-8")
    monkeypatch.setenv("PPT_ACADEMIZER_DECK_KINDS", str(cfg_file))
    patterns = closing_slide_patterns()
    assert len(patterns) == 1 and patterns[0].search("오늘도 수고하셨습니다")
    prs = _deck([("수고하셨습니다", 54, 2.5)])
    cfg = config_for_kind("generic_partner")
    assert is_closing_slide(prs.slides[2], 2, 3, cfg)
    # the built-in phrase no longer matches once replaced
    prs2 = _deck(CLOSING)
    assert not is_closing_slide(prs2.slides[2], 2, 3, cfg)


@pytest.mark.parametrize("text", ["감사합니다", "Thank you!", "Q&A", "Thanks for watching"])
def test_default_phrases(text) -> None:
    assert any(p.search(text) for p in closing_slide_patterns())
