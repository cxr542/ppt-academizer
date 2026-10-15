"""Gray "600 x 400" style pictures are kept but reported; real pictures are not flagged."""

import io

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from core.warnings_display import format_warning
from scripts.build_cmp_academy import is_stock_placeholder_picture, looks_like_placeholder_picture


def _picture(draw, size=(550, 400)):
    im = Image.new("RGB", size, (221, 221, 221))
    draw(ImageDraw.Draw(im), size)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(buf, Inches(1), Inches(1), Inches(4), Inches(3))


def _label(d, size):
    d.text((size[0] // 3, size[1] // 2), "600 x 400", fill=(150, 150, 150))


def test_gray_placeholder_with_a_size_label_is_flagged_but_not_a_dropped_stub() -> None:
    pic = _picture(_label)  # 550x400 is not one of the known stub sizes
    assert not is_stock_placeholder_picture(pic)
    assert looks_like_placeholder_picture(pic)


def test_a_photo_like_picture_is_not_flagged() -> None:
    def noisy(d, size):
        for x in range(0, size[0], 7):
            for y in range(0, size[1], 7):
                d.rectangle((x, y, x + 6, y + 6), fill=((x * 5) % 255, (y * 3) % 255, (x + y) % 255))

    assert not looks_like_placeholder_picture(_picture(noisy))


def test_a_dark_flat_picture_is_not_flagged() -> None:
    im = Image.new("RGB", (550, 400), (40, 60, 120))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    prs = Presentation()
    pic = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(buf, Inches(1), Inches(1), Inches(4), Inches(3))
    assert not looks_like_placeholder_picture(pic)


def test_small_icons_are_not_flagged() -> None:
    assert not looks_like_placeholder_picture(_picture(_label, size=(120, 120)))


def test_warning_message_names_the_count_and_slide() -> None:
    out = format_warning({"code": "PLACEHOLDER_LOOKALIKE", "src_index": 6, "count": 1})
    assert out["level"] == "warn" and out["message"].startswith("원본 7장: ")
    assert "1개" in out["message"] and "그대로 두었습니다" in out["message"] and "처럼 보이는" in out["message"]
