from __future__ import annotations

from base64 import b64decode
from dataclasses import dataclass

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from pptx.util import Pt

from scripts.build_cmp_academy import (
    DOCUMENT_BODY_PT,
    _add_document_textbox,
    reflow_governing_x_body_text,
    apply_academy_body_shape_typography,
    copy_body_shapes_filtered,
    hide_empty_governing_placeholder,
    restore_content_step_placeholder,
    should_drop_body_picture,
    should_preserve_primary_visual,
)
from pptx.enum.text import PP_ALIGN  # noqa: F401  (kept for parity with the local tests)


@dataclass(frozen=True)
class PictureShape:
    left: int
    top: int
    width: int
    height: int
    shape_type: int = MSO_SHAPE_TYPE.PICTURE


@dataclass(frozen=True)
class PlaceholderFormat:
    idx: int


class ParentElement:
    def __init__(self) -> None:
        self.removed = False

    def remove(self, _element) -> None:
        self.removed = True


class ShapeElement:
    def __init__(self, parent: ParentElement) -> None:
        self.parent = parent

    def getparent(self) -> ParentElement:
        return self.parent


class PlaceholderShape:
    def __init__(
        self,
        idx: int,
        text: str = "",
        box: tuple[int, int, int, int] = (0, 0, 0, 0),
    ) -> None:
        self.placeholder_format = PlaceholderFormat(idx)
        self.text = text
        self.left, self.top, self.width, self.height = box
        self.parent = ParentElement()
        self._element = ShapeElement(self.parent)


@dataclass(frozen=True)
class SlideLayout:
    placeholders: list[PlaceholderShape]


@dataclass(frozen=True)
class PlaceholderSlide:
    placeholders: list[PlaceholderShape]
    slide_layout: SlideLayout | None = None


class Font:
    def __init__(self) -> None:
        self.size = None
        self.name = None
        self.color = Color()
        self.bold = None


class Color:
    def __init__(self) -> None:
        self.rgb = None


class Run:
    def __init__(self, text: str) -> None:
        self.text = text
        self.font = Font()


class Paragraph:
    def __init__(self, text: str) -> None:
        self.runs = [Run(text)] if text else []
        self.alignment = None


class TextFrame:
    def __init__(self) -> None:
        self.paragraphs: list[Paragraph] = []
        self.word_wrap = None
        self.auto_size = None

    @property
    def text(self) -> str:
        return "\n".join(run.text for para in self.paragraphs for run in para.runs)

    @text.setter
    def text(self, value: str) -> None:
        self.paragraphs = [Paragraph(line) for line in value.split("\n")]


class TextBox:
    def __init__(self) -> None:
        self.text_frame = TextFrame()
        self.has_text_frame = True
        self.is_placeholder = False
        self.shape_type = MSO_SHAPE_TYPE.TEXT_BOX
        self.top = 1_300_000

    @property
    def text(self) -> str:
        return self.text_frame.text


class Shapes:
    def __init__(self) -> None:
        self.added: list[TextBox] = []

    def add_textbox(self, _left: int, _top: int, _width: int, _height: int) -> TextBox:
        box = TextBox()
        self.added.append(box)
        return box

    def __iter__(self):
        return iter(self.added)


class TextboxSlide:
    def __init__(self) -> None:
        self.shapes = Shapes()


def test_left_body_pictures_are_dropped_for_editable_body() -> None:
    shape = PictureShape(left=762_000, top=1_900_000, width=5_000_000, height=2_000_000)

    assert should_drop_body_picture(
        shape,
        slide_width=12_192_000,
        slide_height=6_858_000,
        permissive_raster=True,
        editable_body=True,
    )


def test_left_body_pictures_are_kept_by_default() -> None:
    """Editable body is opt-in: the default migrate path still copies the body as is."""
    shape = PictureShape(left=762_000, top=1_900_000, width=5_000_000, height=2_000_000)

    assert not should_drop_body_picture(
        shape,
        slide_width=12_192_000,
        slide_height=6_858_000,
        permissive_raster=True,
    )


def test_right_column_pictures_are_kept_as_visual_assets() -> None:
    shape = PictureShape(left=6_300_000, top=1_900_000, width=5_000_000, height=2_000_000)

    assert not should_drop_body_picture(
        shape,
        slide_width=12_192_000,
        slide_height=6_858_000,
        permissive_raster=True,
        editable_body=True,
    )


def test_visual_heavy_slide_preserves_primary_picture_when_text_is_insufficient(
    tmp_path,
) -> None:
    image_path = tmp_path / "visual.png"
    image_path.write_bytes(
        b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAAEklEQVR4nGP4z8AAAAMBAQDJ"
            "/pLvAAAAAElFTkSuQmCC"
        )
    )
    src_prs = Presentation()
    dst_prs = Presentation()
    blank_src = src_prs.slide_layouts[6]
    blank_dst = dst_prs.slide_layouts[6]
    src_slide = src_prs.slides.add_slide(blank_src)
    dst_slide = dst_prs.slides.add_slide(blank_dst)
    pic = src_slide.shapes.add_picture(
        str(image_path),
        Inches(0.5),
        Inches(1.6),
        width=Inches(5.0),
        height=Inches(3.0),
    )
    slide_width = int(src_prs.slide_width)
    slide_height = int(src_prs.slide_height)

    assert should_drop_body_picture(
        pic,
        slide_width=slide_width,
        slide_height=slide_height,
        permissive_raster=True,
        editable_body=True,
    )
    assert should_preserve_primary_visual(
        src_slide,
        title=None,
        governing=None,
        slide_width=slide_width,
        slide_height=slide_height,
    )

    copy_body_shapes_filtered(
        src_slide,
        dst_slide,
        title=None,
        governing=None,
        slide_width=slide_width,
        slide_height=slide_height,
        preserve_primary_visual=True,
        editable_body=True,
    )

    pictures = [
        shape for shape in dst_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1
    assert int(pictures[0].top) >= 1_150_000


def test_empty_governing_placeholder_is_parked_off_slide() -> None:
    # main hides the empty guide by moving it off the slide (the local branch removed it).
    ph13 = PlaceholderShape(13)
    slide = PlaceholderSlide(placeholders=[PlaceholderShape(10, "Title"), ph13])

    hide_empty_governing_placeholder(slide)

    assert ph13.left == 99999999 and ph13.top == 99999999


def test_step_placeholder_returns_to_layout_box_after_polish_expands_it() -> None:
    expected = (515_938, 360_540, 121_828, 215_444)
    ph12 = PlaceholderShape(12, "1", box=(514_984, 757_189, 5_760_404, 411_480))
    slide = PlaceholderSlide(
        placeholders=[ph12],
        slide_layout=SlideLayout(placeholders=[PlaceholderShape(12, "0", box=expected)]),
    )

    restore_content_step_placeholder(slide)

    assert (ph12.left, ph12.top, ph12.width, ph12.height) == expected


def test_reflow_document_textbox_uses_body_shape_font_size() -> None:
    slide = TextboxSlide()

    _add_document_textbox(
        slide,
        left=0,
        top=0,
        width=1_000,
        height=1_000,
        blocks=["개념 정의\n본문", "가치:\n값", "주요 도구:\n도구"],
    )

    runs = [
        run
        for box in slide.shapes.added
        for para in box.text_frame.paragraphs
        for run in para.runs
    ]
    assert runs
    assert all(run.font.size == Pt(DOCUMENT_BODY_PT) for run in runs)


def test_reflow_document_textbox_keeps_body_size_after_typography_polish() -> None:
    slide = TextboxSlide()

    _add_document_textbox(
        slide,
        left=0,
        top=0,
        width=1_000,
        height=1_000,
        blocks=["개념 정의\n본문", "가치:\n값", "주요 도구:\n도구"],
    )
    apply_academy_body_shape_typography(slide)

    runs = [
        run
        for box in slide.shapes.added
        for para in box.text_frame.paragraphs
        for run in para.runs
    ]
    assert runs
    assert all(run.font.size == Pt(DOCUMENT_BODY_PT) for run in runs)


# --- real python-pptx slides: reflow safety (not in the local branch) ---------------------------

from pptx.util import Inches as _In  # noqa: E402


def _body_slide(texts, *, extra=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = _In(13.333), _In(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, text in enumerate(texts):
        slide.shapes.add_textbox(_In(1), _In(1.6 + i * 1.1), _In(8), _In(1)).text_frame.text = text
    if extra:
        extra(slide)
    return prs, slide


def _doc_boxes(slide):
    return [s for s in slide.shapes if s.name == "document_body"]


def test_reflow_merges_unique_text_into_one_document_box() -> None:
    prs, slide = _body_slide(["개념 정의\n본문", "가치:\n값", "개념 정의\n본문"])  # one duplicate
    assert reflow_governing_x_body_text(slide, int(prs.slide_width), int(prs.slide_height))

    boxes = _doc_boxes(slide)
    assert len(boxes) == 1
    text = boxes[0].text_frame.text
    assert "개념 정의" in text and "가치:" in text and text.count("개념 정의") == 1
    # the original text boxes are gone: only the document box remains
    assert [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()] == boxes


def test_reflow_keeps_the_text_left_of_a_right_column_picture(tmp_path) -> None:
    from PIL import Image

    img = tmp_path / "v.png"
    Image.new("RGB", (400, 300), (10, 100, 200)).save(img)
    prs, slide = _body_slide(
        ["개념 정의\n본문", "가치:\n값"],
        extra=lambda s: s.shapes.add_picture(str(img), _In(8.2), _In(1.8), _In(4.5), _In(3.4)),
    )
    assert reflow_governing_x_body_text(slide, int(prs.slide_width), int(prs.slide_height))

    box = _doc_boxes(slide)[0]
    assert int(box.left) + int(box.width) <= int(_In(8.2))  # ends before the picture starts


def test_reflow_refuses_when_a_table_sits_in_the_left_column() -> None:
    prs, slide = _body_slide(
        ["개념 정의\n본문", "가치:\n값"],
        extra=lambda s: s.shapes.add_table(3, 3, _In(0.8), _In(3.5), _In(7), _In(2)),
    )
    before = [s.text_frame.text for s in slide.shapes if s.has_text_frame]

    assert not reflow_governing_x_body_text(slide, int(prs.slide_width), int(prs.slide_height))
    assert _doc_boxes(slide) == []
    assert [s.text_frame.text for s in slide.shapes if s.has_text_frame] == before  # untouched


def test_reflow_with_no_body_text_does_nothing() -> None:
    prs, slide = _body_slide([])
    assert not reflow_governing_x_body_text(slide, int(prs.slide_width), int(prs.slide_height))
