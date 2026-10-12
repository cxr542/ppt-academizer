#!/usr/bin/env python3
"""Tests for Google Slides image-only pptx ingest."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402

from scripts.pptx_ingest import (  # noqa: E402
    is_image_only_slide,
    slide_text_blocks,
    export_slide_background_image,
)

APPLE_FIXTURE = (
    ROOT.parent.parent.parent
    / "apps"
    / "ppt-academizer"
    / "tests"
    / "fixtures"
    / "apple-history.pptx"
)


@pytest.mark.skipif(not APPLE_FIXTURE.is_file(), reason="apple-history.pptx fixture missing")
def test_apple_history_slides_are_image_only():
    prs = Presentation(str(APPLE_FIXTURE))
    assert len(prs.slides) == 7
    for slide in prs.slides:
        assert slide_text_blocks(slide) == []
        assert is_image_only_slide(slide)


@pytest.mark.skipif(not APPLE_FIXTURE.is_file(), reason="apple-history.pptx fixture missing")
def test_export_background(tmp_path: Path):
    prs = Presentation(str(APPLE_FIXTURE))
    dest = tmp_path / "s1.png"
    out = export_slide_background_image(prs.slides[0], dest)
    assert out is not None
    assert out.stat().st_size > 1000
