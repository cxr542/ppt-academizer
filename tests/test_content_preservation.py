"""Silent content loss guards: body placeholders, long governing text, CONTENT_LOSS warning."""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches

from core.warnings_display import format_warning
from scripts.build_cmp_academy import content_loss_warning, slide_text_chars
from scripts.shape_migrate_lib import is_body_shape


def _title_and_content_slide(body: str):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "쿠버네티스 핵심 개념"
    slide.placeholders[1].text_frame.text = body
    return slide


def test_body_placeholder_with_text_is_kept_title_is_not() -> None:
    slide = _title_and_content_slide("Pod는 배포의 최소 단위입니다")
    assert not is_body_shape(slide.shapes.title, "쿠버네티스 핵심 개념", None)
    assert is_body_shape(slide.placeholders[1], "쿠버네티스 핵심 개념", None)


def test_empty_body_placeholder_is_dropped() -> None:
    slide = _title_and_content_slide("")
    assert not is_body_shape(slide.placeholders[1], "쿠버네티스 핵심 개념", None)


def test_body_placeholder_matching_governing_is_still_deduplicated() -> None:
    text = "Pod는 배포의 최소 단위입니다"
    slide = _title_and_content_slide(text)
    # When the same text was used for the governing line it must not be copied twice.
    assert not is_body_shape(slide.placeholders[1], "쿠버네티스 핵심 개념", text)


def test_slide_text_chars_counts_text_frames_and_table_cells() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "가 나 다"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "라 마"
    table.cell(1, 1).text = "바"
    assert slide_text_chars(slide) == 3 + 2 + 1


def test_content_loss_warning_thresholds() -> None:
    assert content_loss_warning(3, 30, 0) is None  # too short to judge
    assert content_loss_warning(3, 100, 70) is None  # exactly 70% is fine
    w = content_loss_warning(3, 338, 14)
    assert w and w["code"] == "CONTENT_LOSS" and w["src_index"] == 3
    shown = format_warning(w)
    assert shown["level"] == "warn" and "4" in shown["message"]  # 1-based slide label


def test_slide_skipped_warning_is_displayed() -> None:
    shown = format_warning({"code": "SLIDE_SKIPPED", "src_index": 8})
    assert shown["level"] == "warn" and "9" in shown["message"]
