"""A content picture that starts in the header band must not cover the academy title row."""

import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import scripts.shape_migrate_lib as sml


@pytest.fixture(autouse=True)
def _blank_is_a_content_layout(monkeypatch):
    monkeypatch.setattr(sml, "CONTENT_LAYOUTS", set(sml.CONTENT_LAYOUTS) | {"Blank"})


def _slide_with_picture(left, top, width, height):
    buf = io.BytesIO()
    Image.new("RGB", (200, 200), (200, 200, 200)).save(buf, "PNG")
    buf.seek(0)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(buf, Inches(left), Inches(top), Inches(width), Inches(height))
    return prs, slide, pic


def test_tall_right_picture_moves_below_the_header_and_keeps_its_right_edge() -> None:
    prs, slide, pic = _slide_with_picture(6.67, 0.5, 5.78, 6.5)  # real deck: VMware slide 11
    right, aspect = pic.left + pic.width, pic.width / pic.height

    assert sml.fit_header_overlapping_pictures(slide, int(prs.slide_width), int(prs.slide_height)) == 1

    assert pic.top >= sml.HEADER_BOTTOM_EMU
    assert pic.top + pic.height <= int(prs.slide_height) - sml.SLIDE_MARGIN_BOTTOM
    assert abs((pic.left + pic.width) - right) <= 1
    assert abs(pic.width / pic.height - aspect) < 0.01


def test_left_picture_keeps_its_left_edge() -> None:
    prs, slide, pic = _slide_with_picture(0.6, 0.4, 4.0, 6.6)
    left = pic.left
    sml.fit_header_overlapping_pictures(slide, int(prs.slide_width), int(prs.slide_height))
    assert pic.left == left and pic.top >= sml.HEADER_BOTTOM_EMU


def test_picture_already_below_the_header_is_untouched() -> None:
    prs, slide, pic = _slide_with_picture(6.67, 1.6, 5.0, 4.0)
    box = (pic.left, pic.top, pic.width, pic.height)
    assert sml.fit_header_overlapping_pictures(slide, int(prs.slide_width), int(prs.slide_height)) == 0
    assert (pic.left, pic.top, pic.width, pic.height) == box


def test_non_content_layouts_are_ignored(monkeypatch) -> None:
    monkeypatch.setattr(sml, "CONTENT_LAYOUTS", {"내지_거버닝 O"})
    prs, slide, pic = _slide_with_picture(6.67, 0.5, 5.78, 6.5)
    assert sml.fit_header_overlapping_pictures(slide, int(prs.slide_width), int(prs.slide_height)) == 0
