"""The template's example course name on 목차/간지 layouts is replaced by the deck's own title."""

import json

from pptx import Presentation
from pptx.util import Inches

import scripts.build_cmp_academy as bca
from scripts.build_cmp_academy import apply_course_name_to_layouts, fit_course_name_pt

EXAMPLE = "Introduction to CLOUD COMPUTING"


class _Layout:
    def __init__(self, name, shapes):
        self.name, self.shapes = name, shapes


class _Deck:
    def __init__(self, layouts):
        self.slide_layouts = layouts


def _prs_with_layout_text(monkeypatch, layout_name="간지", text=EXAMPLE):
    """A deck stub whose given layout carries a real python-pptx text box."""
    slide = Presentation().slides.add_slide(Presentation().slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.66), Inches(0.78), Inches(5.18), Inches(0.4))
    box.text_frame.text = text
    return _Deck([_Layout(layout_name, slide.shapes)]), box


def test_example_text_is_replaced_on_section_and_toc_layouts(monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    prs, box = _prs_with_layout_text(monkeypatch)
    assert apply_course_name_to_layouts(prs, "오케스트로 VMware 윈백 전략") == 1
    assert box.text_frame.text == "오케스트로 VMware 윈백 전략"


def test_other_text_and_other_layouts_are_untouched(monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    prs, box = _prs_with_layout_text(monkeypatch, text="CONTENTS")
    assert apply_course_name_to_layouts(prs, "과정명") == 0 and box.text_frame.text == "CONTENTS"
    prs2, box2 = _prs_with_layout_text(monkeypatch, layout_name="내지_거버닝 O")
    assert apply_course_name_to_layouts(prs2, "과정명") == 0 and box2.text_frame.text == EXAMPLE


def test_no_name_means_no_change(monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    prs, box = _prs_with_layout_text(monkeypatch)
    assert apply_course_name_to_layouts(prs, None) == 0 and apply_course_name_to_layouts(prs, "  ") == 0
    assert box.text_frame.text == EXAMPLE


def test_config_can_disable_or_change_the_example(tmp_path, monkeypatch) -> None:
    cfg = tmp_path / "kinds.json"
    monkeypatch.setenv("PPT_ACADEMIZER_DECK_KINDS", str(cfg))
    cfg.write_text(json.dumps({"template_course_name_text": ""}), encoding="utf-8")
    prs, box = _prs_with_layout_text(monkeypatch)
    assert apply_course_name_to_layouts(prs, "과정명") == 0 and box.text_frame.text == EXAMPLE
    cfg.write_text(json.dumps({"template_course_name_text": "Sample Course"}), encoding="utf-8")
    prs2, box2 = _prs_with_layout_text(monkeypatch, text="Sample Course")
    assert apply_course_name_to_layouts(prs2, "과정명") == 1 and box2.text_frame.text == "과정명"


def test_long_names_get_a_smaller_font() -> None:
    assert fit_course_name_pt("짧은 과정명") == 24
    assert fit_course_name_pt("가" * 30) == 24
    assert fit_course_name_pt("가" * 36) == 18
    assert fit_course_name_pt("가" * 50) == 14


def test_whitespace_from_line_breaks_is_collapsed(monkeypatch) -> None:
    monkeypatch.delenv("PPT_ACADEMIZER_DECK_KINDS", raising=False)
    prs, box = _prs_with_layout_text(monkeypatch)
    apply_course_name_to_layouts(prs, "오케스트로 VMware 윈백 \x0b 시장\n주도 전략")
    assert box.text_frame.text == "오케스트로 VMware 윈백 시장 주도 전략"
