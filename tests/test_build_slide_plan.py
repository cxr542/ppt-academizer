"""Slide plan fidelity: empty slides stay in plan as content."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from scripts.build_cmp_academy import build_slide_plan, expected_output_slide_count
from scripts.deck_migrate_config import config_for_kind


def test_empty_slide_in_plan(tmp_path: Path) -> None:
    src = tmp_path / "empty-slide.pptx"
    prs = Presentation()
    blank = prs.slides.add_slide(prs.slide_layouts[6])
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(src))

    cfg = config_for_kind("contrabass")
    plan = build_slide_plan(prs, cfg)
    assert len(plan) >= 2
    kinds = [k for _, k in plan]
    assert "content" in kinds
    assert expected_output_slide_count(plan, cfg) == len(plan)
