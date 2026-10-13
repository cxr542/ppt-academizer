"""Two-digit step numbers must fit the layout's one-digit-wide number box."""

from pptx import Presentation
from pptx.util import Emu, Inches

from scripts.build_cmp_academy import STEP_NUMBER_EMU_PER_DIGIT, widen_step_number_box


def _slide_with_number(text: str, width=Emu(121_828)):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.56), Inches(0.39), width, Inches(0.24))
    box.text_frame.text = text
    box.text_frame.word_wrap = True
    return slide, box


def test_two_digit_number_gets_a_wider_non_wrapping_box() -> None:
    slide, box = _slide_with_number("10")
    assert widen_step_number_box(slide, "10")
    assert box.width >= STEP_NUMBER_EMU_PER_DIGIT * 2
    assert box.text_frame.word_wrap is False


def test_single_digit_is_left_alone() -> None:
    slide, box = _slide_with_number("7")
    assert not widen_step_number_box(slide, "7")
    assert box.width == 121_828


def test_a_wide_text_box_with_the_same_text_is_not_touched() -> None:
    slide, box = _slide_with_number("10", width=Inches(4))
    assert not widen_step_number_box(slide, "10")
    assert box.width == Inches(4)
