"""Text decks (e.g. apple-history) are analysed as text decks; auto routing is migrate_cmp.

Originally ported from ppt-test, which expected ``spec``; auto routing has been forced to
``migrate_cmp`` (spec drops images/tables/charts), so the expectation follows that.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from scripts.deck_profile import detect_deck_profile


def test_apple_filename_with_text_routes_migrate_and_counts_text(tmp_path: Path) -> None:
    p = tmp_path / "apple-history-test.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    box.text_frame.text = "본문 단락\n두 번째 줄"
    prs.save(str(p))

    prof, meta = detect_deck_profile(p, filename_hint=p.name)
    assert prof == "migrate_cmp"
    assert meta.get("total_text_blocks", 0) >= 1
