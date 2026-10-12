"""§6.7 raster diagram detection."""

from __future__ import annotations

from pptx import Presentation

from scripts.build_cmp_academy import count_diagram_rasters, is_raster_heavy_slide


def test_empty_slide_has_no_diagram_rasters() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    assert count_diagram_rasters(slide, sw, sh) == 0
    assert not is_raster_heavy_slide(slide, sw, sh)
